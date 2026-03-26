from __future__ import annotations

from datetime import date, datetime
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, PieChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from app.services.report_data_builder import (
    EmployeeAggregateRow,
    ManagerReportData,
    ProjectAggregateRow,
    RegistrySheetRow,
    SupplierAggregateRow,
)

MONEY_FORMAT = '#,##0.00'
PERCENT_FORMAT = '0.00%'
DATE_FORMAT = 'DD.MM.YYYY'
DATETIME_FORMAT = 'DD.MM.YYYY HH:MM'
BORDER = Border(
    left=Side(style='thin', color='D9D9D9'),
    right=Side(style='thin', color='D9D9D9'),
    top=Side(style='thin', color='D9D9D9'),
    bottom=Side(style='thin', color='D9D9D9'),
)


class ManagerExcelReportBuilder:
    def __init__(self) -> None:
        self.title_fill = PatternFill('solid', fgColor='1F4E78')
        self.header_fill = PatternFill('solid', fgColor='D9EAF7')
        self.accent_fill = PatternFill('solid', fgColor='EAF4EA')
        self.warning_fill = PatternFill('solid', fgColor='FCE4D6')
        self.neutral_fill = PatternFill('solid', fgColor='F3F6FA')
        self.title_font = Font(color='FFFFFF', bold=True, size=16)
        self.header_font = Font(bold=True)
        self.metric_title_font = Font(bold=True, color='1F1F1F')
        self.metric_value_font = Font(bold=True, size=14, color='1F1F1F')
        self.sheet_title_font = Font(color='FFFFFF', bold=True, size=13)

    def build(self, report: ManagerReportData, output_path: Path) -> None:
        workbook = Workbook()
        dashboard = workbook.active
        dashboard.title = 'Дашборд'
        self._build_dashboard_sheet(dashboard, report)
        self._build_projects_sheet(workbook.create_sheet('По проектам'), report)
        self._build_employees_sheet(workbook.create_sheet('По сотрудникам'), report)
        self._build_duplicates_sheet(workbook.create_sheet('Дубли'), report)
        self._build_registry_sheet(workbook.create_sheet('Общий реестр'), report)
        workbook.save(output_path)

    def _build_dashboard_sheet(self, sheet, report: ManagerReportData) -> None:
        sheet.merge_cells('A1:L1')
        title_cell = sheet['A1']
        title_cell.value = 'Управленческий отчет по затратам'
        title_cell.fill = self.title_fill
        title_cell.font = self.title_font
        title_cell.alignment = Alignment(horizontal='center', vertical='center')
        sheet.row_dimensions[1].height = 26

        meta_rows = [
            ('Компания', report.company_name, 'Режим', self._mode_label(report.mode)),
            ('Период', report.period_label, 'Фильтр', report.filter_caption),
            ('С даты', self._display_date(report.date_from), 'По дату', self._display_date(report.date_to)),
            ('Сформирован', self._display_datetime(report.generated_at), '', ''),
        ]
        row_index = 2
        for left_label, left_value, right_label, right_value in meta_rows:
            sheet.cell(row=row_index, column=1, value=left_label).font = self.header_font
            sheet.cell(row=row_index, column=2, value=left_value)
            if right_label:
                sheet.cell(row=row_index, column=5, value=right_label).font = self.header_font
            if right_value:
                sheet.cell(row=row_index, column=6, value=right_value)
            row_index += 1

        metric_specs = [
            ('Всего документов', report.kpis['documents'], self.neutral_fill),
            ('Общая сумма затрат', report.kpis['total_amount'], self.accent_fill),
            ('Сумма НДС', report.kpis['vat_total_amount'], self.accent_fill),
            ('Сумма без дублей', report.kpis['sum_without_duplicates'], self.accent_fill),
            ('Средняя сумма документа', report.kpis['average_amount'], self.neutral_fill),
            ('Число проектов', report.kpis['projects'], self.neutral_fill),
            ('Число сотрудников', report.kpis['employees'], self.neutral_fill),
            ('Документов с НДС', report.kpis['documents_with_vat'], self.neutral_fill),
            ('Сумма документов с НДС', report.kpis['amount_with_vat'], self.accent_fill),
            ('Число поставщиков', report.kpis['suppliers'], self.neutral_fill),
            ('Точные дубли', report.kpis['exact_duplicates'], self.warning_fill),
            ('Вероятные дубли', report.kpis['probable_duplicates'], self.warning_fill),
            ('Доля дублей', report.kpis['duplicate_share'], self.warning_fill),
        ]
        self._write_kpi_grid(sheet, metric_specs, start_row=7)

        top_projects = report.top_projects or [
            ProjectAggregateRow('Нет данных', 0, Decimal('0'), Decimal('0'), Decimal('0'), Decimal('0'), 0, 0, 0, 0, 0.0, None, None)
        ]
        top_employees = report.top_employees or [
            EmployeeAggregateRow('Нет данных', None, 'Не найден', 0, Decimal('0'), Decimal('0'), Decimal('0'), Decimal('0'), 0, 0, 0, 0, 0.0, None, None)
        ]
        top_suppliers = report.top_suppliers or [SupplierAggregateRow('Нет данных', 0, Decimal('0'), Decimal('0'))]

        project_table = self._write_dashboard_table(
            sheet,
            title='Топ проектов',
            start_row=19,
            start_col=1,
            headers=['Проект', 'Сумма', 'НДС', 'Документов'],
            rows=[[row.project_name, row.total_amount, row.vat_total_amount, row.document_count] for row in top_projects],
            money_columns={2, 3},
        )
        employee_table = self._write_dashboard_table(
            sheet,
            title='Топ сотрудников',
            start_row=19,
            start_col=5,
            headers=['Сотрудник', 'Сумма', 'НДС', 'Документов'],
            rows=[[row.employee_name, row.total_amount, row.vat_total_amount, row.document_count] for row in top_employees],
            money_columns={2, 3},
        )
        supplier_table = self._write_dashboard_table(
            sheet,
            title='Топ поставщиков',
            start_row=19,
            start_col=9,
            headers=['Поставщик', 'Сумма', 'НДС', 'Документов'],
            rows=[[row.supplier_name, row.total_amount, row.vat_total_amount, row.document_count] for row in top_suppliers],
            money_columns={2, 3},
        )

        duplicate_control = [
            ['Точные дубли', report.kpis['exact_duplicates']],
            ['Вероятные дубли', report.kpis['probable_duplicates']],
            ['Доля дублей', report.kpis['duplicate_share']],
            ['Сумма без дублей', float(report.kpis['sum_without_duplicates'])],
        ]
        duplicate_control_table = self._write_dashboard_table(
            sheet,
            title='Контроль дублей',
            start_row=29,
            start_col=1,
            headers=['Показатель', 'Значение'],
            rows=duplicate_control,
            percent_columns={2},
            percent_row_indexes={3},
        )
        amount_row = duplicate_control_table['header_row'] + 4
        sheet.cell(row=amount_row, column=2).number_format = MONEY_FORMAT

        largest_rows = [
            [
                row.document_id,
                row.project_name,
                row.employee_name,
                row.vendor or '',
                row.total_amount,
                row.vat_total_amount,
                row.document_date,
                row.created_at,
                row.duplicate_status,
            ]
            for row in report.largest_documents
        ] or [['', 'Нет данных', '', '', Decimal('0'), Decimal('0'), None, None, '']]
        largest_table = self._write_dashboard_table(
            sheet,
            title='Крупнейшие документы',
            start_row=29,
            start_col=5,
            headers=['ID', 'Проект', 'Кто внес', 'Поставщик', 'Сумма', 'НДС', 'Дата документа', 'Дата ввода', 'Статус дубля'],
            rows=largest_rows,
            money_columns={5, 6},
            date_columns={7},
            datetime_columns={8},
        )

        project_chart = BarChart()
        project_chart.title = 'Расходы по проектам'
        project_chart.y_axis.title = 'Сумма'
        project_chart.x_axis.title = 'Проекты'
        project_chart.height = 7
        project_chart.width = 10
        project_data = Reference(sheet, min_col=2, min_row=project_table['header_row'], max_row=project_table['last_row'])
        project_cats = Reference(sheet, min_col=1, min_row=project_table['header_row'] + 1, max_row=project_table['last_row'])
        project_chart.add_data(project_data, titles_from_data=True)
        project_chart.set_categories(project_cats)
        sheet.add_chart(project_chart, 'M19')

        employee_chart = BarChart()
        employee_chart.title = 'Расходы по сотрудникам'
        employee_chart.y_axis.title = 'Сумма'
        employee_chart.x_axis.title = 'Сотрудники'
        employee_chart.height = 7
        employee_chart.width = 10
        employee_data = Reference(sheet, min_col=6, min_row=employee_table['header_row'], max_row=employee_table['last_row'])
        employee_cats = Reference(sheet, min_col=5, min_row=employee_table['header_row'] + 1, max_row=employee_table['last_row'])
        employee_chart.add_data(employee_data, titles_from_data=True)
        employee_chart.set_categories(employee_cats)
        sheet.add_chart(employee_chart, 'M34')

        pie_chart = PieChart()
        pie_chart.title = 'Структура дублей'
        pie_chart.height = 7
        pie_chart.width = 8
        pie_data = Reference(
            sheet,
            min_col=2,
            min_row=duplicate_control_table['header_row'] + 1,
            max_row=duplicate_control_table['header_row'] + 2,
        )
        pie_labels = Reference(
            sheet,
            min_col=1,
            min_row=duplicate_control_table['header_row'] + 1,
            max_row=duplicate_control_table['header_row'] + 2,
        )
        pie_chart.add_data(pie_data, titles_from_data=False)
        pie_chart.set_categories(pie_labels)
        sheet.add_chart(pie_chart, 'M49')

        sheet.freeze_panes = None
        sheet.auto_filter.ref = f"A{project_table['header_row']}:D{project_table['last_row']}"
        self._set_dashboard_widths(sheet)

    def _build_projects_sheet(self, sheet, report: ManagerReportData) -> None:
        headers = [
            'Проект',
            'Количество документов',
            'Общая сумма',
            'НДС',
            'Сумма без дублей',
            'Средняя сумма документа',
            'Количество сотрудников',
            'Количество поставщиков',
            'Точные дубли',
            'Вероятные дубли',
            'Доля дублей, %',
            'Первая дата документа',
            'Последняя дата документа',
        ]
        rows = [
            [
                row.project_name,
                row.document_count,
                row.total_amount,
                row.vat_total_amount,
                row.non_duplicate_amount,
                row.average_amount,
                row.employee_count,
                row.supplier_count,
                row.exact_duplicate_count,
                row.probable_duplicate_count,
                row.duplicate_share,
                row.first_document_date,
                row.last_document_date,
            ]
            for row in report.project_rows
        ]
        self._build_table_sheet(
            sheet,
            title='Отчет по проектам',
            headers=headers,
            rows=rows,
            money_columns={3, 4, 5, 6},
            percent_columns={11},
            date_columns={12, 13},
            max_width=28,
        )
        self._set_column_widths(sheet, {
            1: 28, 2: 18, 3: 16, 4: 14, 5: 18, 6: 18,
            7: 18, 8: 20, 9: 14, 10: 18, 11: 14, 12: 16, 13: 18,
        })

    def _build_employees_sheet(self, sheet, report: ManagerReportData) -> None:
        headers = [
            'Сотрудник',
            'Username',
            'Статус участника',
            'Количество документов',
            'Общая сумма',
            'НДС',
            'Сумма без дублей',
            'Средняя сумма документа',
            'Количество проектов',
            'Количество поставщиков',
            'Точные дубли',
            'Вероятные дубли',
            'Доля дублей, %',
            'Первая дата документа',
            'Последняя дата документа',
        ]
        rows = [
            [
                row.employee_name,
                row.username or '',
                row.member_status,
                row.document_count,
                row.total_amount,
                row.vat_total_amount,
                row.non_duplicate_amount,
                row.average_amount,
                row.project_count,
                row.supplier_count,
                row.exact_duplicate_count,
                row.probable_duplicate_count,
                row.duplicate_share,
                row.first_document_date,
                row.last_document_date,
            ]
            for row in report.employee_rows
        ]
        self._build_table_sheet(
            sheet,
            title='Отчет по сотрудникам',
            headers=headers,
            rows=rows,
            money_columns={5, 6, 7, 8},
            percent_columns={13},
            date_columns={14, 15},
            max_width=26,
        )
        self._set_column_widths(sheet, {
            1: 24, 2: 16, 3: 18, 4: 18, 5: 16, 6: 14, 7: 18, 8: 18,
            9: 18, 10: 20, 11: 14, 12: 18, 13: 14, 14: 16, 15: 18,
        })

    def _build_duplicates_sheet(self, sheet, report: ManagerReportData) -> None:
        headers = [
            'ID документа',
            'ID исходного документа',
            'Статус дубля',
            'Дата ввода',
            'Дата документа',
            'Проект',
            'Кто внес',
            'Username',
            'Поставщик',
            'ИНН поставщика',
            'Номер документа',
            'Сумма',
            'НДС',
        ]
        rows = [
            [
                row.document_id,
                row.source_document_id,
                row.duplicate_status,
                row.created_at,
                row.document_date,
                row.project_name,
                row.uploaded_by_name,
                row.username or '',
                row.vendor or '',
                row.vendor_inn or '',
                row.document_number or '',
                row.total_amount,
                row.vat_total_amount,
            ]
            for row in report.duplicate_rows
        ]
        self._build_table_sheet(
            sheet,
            title='Дубли документов',
            headers=headers,
            rows=rows,
            money_columns={12, 13},
            date_columns={5},
            datetime_columns={4},
            max_width=32,
            wrap_text=False,
            data_row_height=18,
        )
        self._set_column_widths(sheet, {
            1: 12, 2: 14, 3: 18, 4: 18, 5: 16, 6: 20, 7: 20, 8: 14,
            9: 20, 10: 14, 11: 16, 12: 14, 13: 14,
        })

    def _build_registry_sheet(self, sheet, report: ManagerReportData) -> None:
        headers = [
            'ID документа',
            'Дата ввода',
            'Дата документа',
            'Компания',
            'Проект',
            'ID проекта',
            'Сотрудник',
            'Username',
            'Статус участника',
            'Тип документа',
            'Номер документа',
            'Входящий номер',
            'Поставщик',
            'ИНН поставщика',
            'КПП поставщика',
            'Сумма документа',
            'НДС по документу',
            'Тип НДС',
            'Статус дубля',
            'ID исходного дубля',
            'ID позиции',
            '№ строки позиции',
            'Наименование позиции',
            'Количество',
            'Цена',
            'Сумма позиции',
            'Количество позиций в документе',
        ]
        rows = [
            [
                row.document_id,
                row.created_at,
                row.document_date,
                row.company_name,
                row.project_name,
                row.project_id,
                row.employee_name,
                row.username or '',
                row.member_status,
                row.document_type,
                row.document_number or '',
                row.incoming_number or '',
                row.vendor or '',
                row.vendor_inn or '',
                row.vendor_kpp or '',
                row.total_amount,
                row.vat_total_amount,
                self._vat_scope_label(row.vat_scope),
                row.duplicate_status,
                row.source_duplicate_id,
                row.item_id,
                row.item_line_no,
                row.item_name or '',
                row.item_quantity,
                row.item_price,
                row.item_total,
                row.item_count,
            ]
            for row in report.registry_rows
        ]
        self._build_table_sheet(
            sheet,
            title='Реестр всех документов и позиций',
            headers=headers,
            rows=rows,
            money_columns={16, 17, 25, 26},
            date_columns={3},
            datetime_columns={2},
            max_width=36,
            wrap_text=False,
            data_row_height=18,
        )
        self._set_column_widths(sheet, {
            1: 12, 2: 18, 3: 16, 4: 20, 5: 20, 6: 12, 7: 20, 8: 14,
            9: 18, 10: 18, 11: 16, 12: 16, 13: 20, 14: 14, 15: 14, 16: 14,
            17: 14, 18: 16, 19: 16, 20: 14, 21: 12, 22: 14, 23: 26, 24: 12,
            25: 12, 26: 14, 27: 18,
        })

    def _build_table_sheet(
        self,
        sheet,
        *,
        title: str,
        headers: list[str],
        rows: list[list[object]],
        money_columns: set[int] | None = None,
        percent_columns: set[int] | None = None,
        date_columns: set[int] | None = None,
        datetime_columns: set[int] | None = None,
        max_width: int = 50,
        wrap_text: bool = True,
        data_row_height: float | None = None,
    ) -> None:
        money_columns = money_columns or set()
        percent_columns = percent_columns or set()
        date_columns = date_columns or set()
        datetime_columns = datetime_columns or set()

        sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
        title_cell = sheet.cell(row=1, column=1, value=title)
        title_cell.font = self.sheet_title_font
        title_cell.fill = self.title_fill
        title_cell.alignment = Alignment(horizontal='center')
        title_cell.border = BORDER

        for index, header in enumerate(headers, start=1):
            cell = sheet.cell(row=2, column=index, value=header)
            cell.font = self.header_font
            cell.fill = self.header_fill
            cell.border = BORDER
            cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)

        if not rows:
            rows = [['' for _ in headers]]

        for row_index, row_values in enumerate(rows, start=3):
            if data_row_height is not None:
                sheet.row_dimensions[row_index].height = data_row_height
            for col_index, raw_value in enumerate(row_values, start=1):
                cell = sheet.cell(row=row_index, column=col_index, value=self._excel_value(raw_value))
                cell.border = BORDER
                cell.alignment = Alignment(vertical='top', wrap_text=wrap_text)
                if col_index in money_columns and isinstance(raw_value, (Decimal, int, float)):
                    cell.number_format = MONEY_FORMAT
                elif col_index in percent_columns and isinstance(raw_value, (int, float, Decimal)):
                    cell.number_format = PERCENT_FORMAT
                elif col_index in date_columns and raw_value:
                    cell.number_format = DATE_FORMAT
                elif col_index in datetime_columns and raw_value:
                    cell.number_format = DATETIME_FORMAT

        sheet.freeze_panes = 'A3'
        sheet.auto_filter.ref = f"A2:{get_column_letter(len(headers))}{sheet.max_row}"
        self._autofit_widths(sheet, max_width=max_width)

    def _write_kpi_grid(self, sheet, metric_specs: list[tuple[str, object, PatternFill]], start_row: int) -> None:
        card_columns = [1, 4, 7, 10]
        for index, (title, value, fill) in enumerate(metric_specs):
            row_offset = (index // 4) * 3
            col = card_columns[index % 4]
            row = start_row + row_offset
            sheet.merge_cells(start_row=row, start_column=col, end_row=row, end_column=col + 2)
            sheet.merge_cells(start_row=row + 1, start_column=col, end_row=row + 1, end_column=col + 2)
            title_cell = sheet.cell(row=row, column=col, value=title)
            value_cell = sheet.cell(row=row + 1, column=col, value=self._metric_value(title, value))
            for target in (title_cell, value_cell):
                target.fill = fill
                target.border = BORDER
                target.alignment = Alignment(horizontal='center', vertical='center')
            title_cell.font = self.metric_title_font
            value_cell.font = self.metric_value_font
            if title == 'Доля дублей':
                value_cell.number_format = PERCENT_FORMAT
            elif title in {'Общая сумма затрат', 'Сумма НДС', 'Сумма документов с НДС', 'Сумма без дублей', 'Средняя сумма документа'}:
                value_cell.number_format = MONEY_FORMAT
            sheet.row_dimensions[row].height = 20
            sheet.row_dimensions[row + 1].height = 24

    def _write_dashboard_table(
        self,
        sheet,
        *,
        title: str,
        start_row: int,
        start_col: int,
        headers: list[str],
        rows: list[list[object]],
        money_columns: set[int] | None = None,
        percent_columns: set[int] | None = None,
        percent_row_indexes: set[int] | None = None,
        date_columns: set[int] | None = None,
        datetime_columns: set[int] | None = None,
    ) -> dict[str, int]:
        money_columns = money_columns or set()
        percent_columns = percent_columns or set()
        percent_row_indexes = percent_row_indexes or set()
        date_columns = date_columns or set()
        datetime_columns = datetime_columns or set()
        end_col = start_col + len(headers) - 1
        sheet.merge_cells(start_row=start_row, start_column=start_col, end_row=start_row, end_column=end_col)
        title_cell = sheet.cell(row=start_row, column=start_col, value=title)
        title_cell.font = self.sheet_title_font
        title_cell.fill = self.title_fill
        title_cell.alignment = Alignment(horizontal='center')
        title_cell.border = BORDER
        header_row = start_row + 1
        for offset, header in enumerate(headers):
            cell = sheet.cell(row=header_row, column=start_col + offset, value=header)
            cell.font = self.header_font
            cell.fill = self.header_fill
            cell.border = BORDER
            cell.alignment = Alignment(horizontal='center', wrap_text=True)
        current_row = header_row + 1
        for row_index, values in enumerate(rows, start=1):
            for offset, raw_value in enumerate(values, start=1):
                cell = sheet.cell(row=current_row, column=start_col + offset - 1, value=self._excel_value(raw_value))
                cell.border = BORDER
                cell.alignment = Alignment(vertical='top', wrap_text=True)
                relative_col = offset
                if relative_col in money_columns and isinstance(raw_value, (Decimal, int, float)):
                    cell.number_format = MONEY_FORMAT
                elif relative_col in date_columns and raw_value:
                    cell.number_format = DATE_FORMAT
                elif relative_col in datetime_columns and raw_value:
                    cell.number_format = DATETIME_FORMAT
                elif relative_col in percent_columns and (row_index in percent_row_indexes or not percent_row_indexes):
                    cell.number_format = PERCENT_FORMAT
            current_row += 1
        return {'header_row': header_row, 'last_row': max(header_row + 1, current_row - 1)}

    def _metric_value(self, title: str, value: object) -> object:
        if isinstance(value, Decimal):
            return float(value)
        return value

    def _set_dashboard_widths(self, sheet) -> None:
        widths = {
            'A': 18, 'B': 16, 'C': 16,
            'D': 18, 'E': 16, 'F': 16,
            'G': 18, 'H': 16, 'I': 16,
            'J': 18, 'K': 16, 'L': 16,
            'M': 14, 'N': 14, 'O': 14, 'P': 14,
            'Q': 14, 'R': 14,
        }
        for column, width in widths.items():
            sheet.column_dimensions[column].width = width

    def _autofit_widths(self, sheet, *, max_width: int) -> None:
        for column_cells in sheet.columns:
            first_cell = column_cells[0]
            if not hasattr(first_cell, 'column_letter'):
                continue
            column_letter = first_cell.column_letter
            max_length = 0
            for cell in column_cells:
                value = '' if cell.value is None else str(cell.value)
                if len(value) > max_length:
                    max_length = len(value)
            sheet.column_dimensions[column_letter].width = min(max(max_length + 2, 12), max_width)


    def _set_column_widths(self, sheet, widths: dict[int, int]) -> None:
        for column_index, width in widths.items():
            sheet.column_dimensions[get_column_letter(column_index)].width = width

    def _excel_value(self, value: object) -> object:
        if isinstance(value, Decimal):
            return float(value)
        if isinstance(value, datetime):
            return value.replace(tzinfo=None) if value.tzinfo is not None else value
        if isinstance(value, date):
            return value
        return value

    def _mode_label(self, mode: str) -> str:
        return 'За все время' if mode == 'all_time' else 'Период'

    def _vat_scope_label(self, scope: str | None) -> str:
        return {
            'document': 'весь документ',
            'mixed': 'смешанный',
            'no_vat': 'без НДС',
            'unknown': 'не определен',
            None: 'не определен',
            '': 'не определен',
        }.get(scope, scope or 'не определен')

    def _display_date(self, value: date | None) -> str:
        if value is None:
            return '—'
        return value.strftime('%d.%m.%Y')

    def _display_datetime(self, value: datetime) -> str:
        safe_value = value.replace(tzinfo=None) if value.tzinfo is not None else value
        return safe_value.strftime('%d.%m.%Y %H:%M')

    def _clip(self, value: str | None, limit: int) -> str:
        if not value:
            return ''
        if len(value) <= limit:
            return value
        return value[: limit - 1] + '…'
