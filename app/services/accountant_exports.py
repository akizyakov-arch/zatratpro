"""Company-scoped accountant package, independent of messenger/file SDKs."""
import asyncio
from datetime import date, datetime, timedelta, timezone
from io import BytesIO
from pathlib import Path
import tempfile
from weakref import WeakKeyDictionary
from zipfile import ZipFile, ZIP_DEFLATED

from openpyxl import Workbook
from app.config import TMP_DIR
from app.services.async_work import run_blocking
from app.schemas.fiscal import fiscalization_label
from app.services.database import get_pool
from app.services.advances import AdvanceError
from app.services.reporting import ReportService
from app.services.report_exports import _append_row
from app.services.report_periods import REPORT_PERIOD_LABELS, report_period_start

PERIODS = dict(REPORT_PERIOD_LABELS, previous_month='Прошлый месяц')
MAX_SOURCE_BYTES = 100 * 1024 * 1024
MAX_ARCHIVE_BYTES = 45 * 1024 * 1024
MAX_DOCUMENTS = 2000
ORIGINAL_BATCH_BYTES = 8 * 1024 * 1024
ORIGINAL_BATCH_COUNT = 128
_EXPORT_SLOTS = WeakKeyDictionary()


def _export_slots():
    loop = asyncio.get_running_loop()
    return _EXPORT_SLOTS.setdefault(loop, asyncio.Semaphore(2))


async def _file_work(function, *args):
    return await run_blocking(function, *args)


def _original_batches(rows):
    batch, size = [], 0
    for row in rows:
        amount = row['original_size'] or 0
        if not amount:
            continue
        if batch and (size + amount > ORIGINAL_BATCH_BYTES or len(batch) >= ORIGINAL_BATCH_COUNT):
            yield batch
            batch, size = [], 0
        batch.append(row['id'])
        size += amount
    if batch:
        yield batch


def _spool_originals(rows, directory):
    paths = {}
    for row in rows:
        path = Path(directory) / str(row['document_id'])
        path.write_bytes(row['content'])
        paths[row['document_id']] = path
    return paths


def accountant_dates(value, *, now=None):
    """Inclusive calendar dates, UTC. Custom range: YYYY-MM-DD YYYY-MM-DD."""
    today = (now or datetime.now(timezone.utc)).astimezone(timezone.utc).date()
    if value == 'previous_month':
        end = today.replace(day=1) - timedelta(days=1)
        return end.replace(day=1), end
    if value in PERIODS:
        start = date.min if value == 'all_time' else report_period_start(value, now=now).date()
        return start, today
    try:
        left, right = value.split()
        start, end = date.fromisoformat(left), date.fromisoformat(right)
        if not start <= end or end == date.max:
            raise ValueError
        return start, end
    except (ValueError, TypeError):
        raise AdvanceError('Укажите даты: /accountant 2026-09-01 2026-10-31. Обе даты включены.') from None


class AccountantExportService:
    async def authorize(self, actor):
        async with get_pool().acquire() as c:
            return await ReportService()._manager(c, actor)

    async def build(self, actor, value, *, project_id=None, employee_id=None):
        # Bound concurrent packaging and its unavoidable final <=45MiB payloads.
        async with _export_slots():
            return await self._build(actor, value, project_id=project_id, employee_id=employee_id)

    async def _build(self, actor, value, *, project_id=None, employee_id=None):
        start, end = accountant_dates(value)
        TMP_DIR.mkdir(parents=True, exist_ok=True)
        with tempfile.TemporaryDirectory(prefix='accountant-', dir=TMP_DIR) as directory:
            async with get_pool().acquire() as c, c.transaction(isolation='repeatable_read', readonly=True):
                company = await ReportService()._manager(c, actor)
                await ReportService._scope(c, company, project_id, employee_id)
                rows = await c.fetch('''SELECT d.id,d.document_type,d.external_document_number,
                    d.document_date,d.created_at,d.vendor,d.vendor_inn,d.vendor_kpp,d.currency,
                    d.total_amount,d.duplicate_status,d.is_fiscalized,p.name AS project_name,
                    concat_ws(' ',u.first_name,u.last_name) AS employee_name,
                    o.extension,octet_length(o.content) AS original_size
                    FROM documents d JOIN projects p ON p.id=d.project_id
                    JOIN users u ON u.id=d.uploaded_by_user_id
                    LEFT JOIN document_originals o ON o.document_id=d.id
                    WHERE d.company_id=$1
                      AND coalesce((d.document_date AT TIME ZONE 'UTC')::date,(d.created_at AT TIME ZONE 'UTC')::date) BETWEEN $2 AND $3
                      AND ($4::bigint IS NULL OR d.project_id=$4)
                      AND ($5::bigint IS NULL OR d.uploaded_by_user_id=$5)
                    ORDER BY coalesce((d.document_date AT TIME ZONE 'UTC')::date,(d.created_at AT TIME ZONE 'UTC')::date),d.id
                    LIMIT $6''', company, start, end, project_id, employee_id, MAX_DOCUMENTS + 1)
                if not rows:
                    raise AdvanceError('За выбранный период документов нет.')
                if len(rows) > MAX_DOCUMENTS or sum(r['original_size'] or 0 for r in rows) > MAX_SOURCE_BYTES:
                    raise AdvanceError('Выгрузка слишком большая. Выберите более короткий период.')
                originals = {}
                for ids in _original_batches(rows):
                    batch = await c.fetch('SELECT document_id,content FROM document_originals WHERE document_id=ANY($1::bigint[])', ids)
                    originals.update(await _file_work(_spool_originals, batch, directory))
                    del batch
                items = await c.fetch('SELECT * FROM document_items WHERE document_id=ANY($1::bigint[]) ORDER BY document_id,line_no', [r['id'] for r in rows])
            payload, missing = await _file_work(_package, rows, items, originals, start, end)
            # Access may have changed while constructing a large package.
            if await self.authorize(actor) != company:
                raise AdvanceError('Компания изменилась. Сформируйте выгрузку заново.')
            return f'accountant-{start:%Y%m%d}-{end:%Y%m%d}.zip', payload, missing


def _package(rows, items, originals, start, end):
    workbook = Workbook()
    summary = workbook.active
    summary.title = 'Период'
    _append_row(summary, ['С', start.isoformat(), 'По (включительно)', end.isoformat()])
    _append_row(summary, ['Отбор', 'Дата документа; если не указана — дата загрузки UTC'])
    _append_row(summary, ['Дубли', 'Включены и отмечены; архивные проекты включены'])
    _append_row(summary, ['Фискализация', 'По реквизитам OCR; проверка через ФНС не выполнялась'])
    missing = sum(r['id'] not in originals for r in rows)
    _append_row(summary, ['Документов', len(rows), 'Без скана', missing])
    sheet = workbook.create_sheet('Документы')
    _append_row(sheet, ['ID', 'Проект', 'Сотрудник', 'Тип', 'Номер', 'Дата документа', 'Загружен UTC', 'Поставщик', 'ИНН', 'КПП', 'Валюта', 'Сумма', 'Дубль', 'Файл', 'Основание периода', 'Фискализация'])
    with tempfile.SpooledTemporaryFile(max_size=1024 * 1024, mode='w+b', dir=TMP_DIR) as output:
        return _write_package(output, workbook, rows, items, originals, sheet, missing)


def _write_package(output, workbook, rows, items, originals, sheet, missing):
    with ZipFile(output, 'w', ZIP_DEFLATED) as archive:
        for row in rows:
            path = f"files/{row['id']}.{row['extension']}" if row['id'] in originals else ''
            _append_row(sheet, [row['id'], row['project_name'], row['employee_name'], row['document_type'], row['external_document_number'] or '', str(row['document_date'] or ''), row['created_at'].isoformat(), row['vendor'] or '', row['vendor_inn'] or '', row['vendor_kpp'] or '', row['currency'], str(row['total_amount']), row['duplicate_status'], path or 'Скан отсутствует', 'Дата документа' if row['document_date'] else 'Дата загрузки UTC', fiscalization_label(row['is_fiscalized'])])
            if path:
                original = originals[row['id']]
                if isinstance(original, (bytes, bytearray)):
                    archive.writestr(path, original)
                else:
                    archive.write(original, arcname=path)
                if output.tell() > MAX_ARCHIVE_BYTES:
                    raise AdvanceError('Архив превышает 45 МБ. Выберите более короткий период.')
        positions = workbook.create_sheet('Позиции')
        _append_row(positions, ['ID документа', 'Строка', 'Наименование', 'Количество', 'Цена', 'Сумма'])
        for item in items:
            _append_row(positions, [item['document_id'], item['line_no'], item['name'], _number_text(item['quantity']), _number_text(item['price']), _number_text(item['line_total'])])
        data = BytesIO()
        for worksheet in workbook:
            for cells in worksheet:
                for cell in cells:
                    if isinstance(cell.value, str):
                        cell.data_type = 's'
        workbook.save(data)
        archive.writestr('Реестр.xlsx', data.getvalue())
    if output.tell() > MAX_ARCHIVE_BYTES:
        raise AdvanceError('Архив превышает 45 МБ. Выберите более короткий период.')
    output.seek(0)
    return output.read(), missing


def _number_text(value):
    return '' if value is None else str(value)
