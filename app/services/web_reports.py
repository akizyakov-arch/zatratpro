"""Canonical-user document views and complete reports shared by web clients.

Dates select document upload time and financial confirmation time respectively.
Amounts of different currencies are never combined into the RUB total.
"""
from collections import defaultdict
from datetime import datetime, timezone
from decimal import Decimal
from io import BytesIO
from types import SimpleNamespace

from openpyxl import load_workbook

from app.services.advances import AdvanceError, AdvanceService
from app.services.business_cards import BusinessCardService
from app.services.database import get_pool
from app.services.money_reports import CASH_FIELDS, cash_and_employee_debt
from app.services.report_exports import build_manager_report_workbook, _append_row, _ILLEGAL_XML


_DOCUMENT_SELECT = """
    SELECT d.*, p.name AS project_name, p.status AS project_status,
           COALESCE(NULLIF(d.external_document_number,''),NULLIF(d.incoming_number,'')) AS document_number,
           COALESCE(NULLIF(concat_ws(' ',u.first_name,u.last_name),''),u.username,'Сотрудник №'||u.id) AS uploaded_by_name,
           u.username AS uploaded_by_username,
           CASE WHEN ae.id IS NOT NULL THEN 'advance'
                WHEN ce.id IS NOT NULL THEN 'card' ELSE NULL END AS payment_source,
           ae.advance_id, ce.card_id,
           COALESCE(ae.approved_at,ce.approved_at) AS expense_approved_at,
           COALESCE(ae.id,ce.id) AS funding_expense_id
    FROM documents d
    JOIN projects p ON p.id=d.project_id AND p.company_id=d.company_id
    JOIN users u ON u.id=d.uploaded_by_user_id
    LEFT JOIN advance_expenses ae ON ae.document_id=d.id AND ae.rejected_at IS NULL
    LEFT JOIN card_expenses ce ON ce.document_id=d.id AND ce.rejected_at IS NULL
"""


def _period(start, end):
    if (not isinstance(start, datetime) or not isinstance(end, datetime)
            or start.tzinfo is None or end.tzinfo is None or start >= end):
        raise AdvanceError('Укажите корректный период с часовым поясом.')


def _page(page, per_page):
    if (isinstance(page, bool) or isinstance(per_page, bool)
            or not isinstance(page, int) or not isinstance(per_page, int)
            or page < 1 or not 1 <= per_page <= 100):
        raise AdvanceError('Некорректная страница списка.')


class WebReportService:
    async def _manager(self, connection, actor, *, lock=False):
        access = AdvanceService()
        member = await (access._lock_member(connection, actor) if lock else access._member(connection, actor))
        if member['role'] != 'manager':
            raise AdvanceError('Панель документов доступна руководителю компании.')
        return member['company_id']

    @staticmethod
    async def _scope(connection, company, project_id, employee_id):
        # Historical participants remain selectable even after removal.
        if project_id is not None and not await connection.fetchval(
                'SELECT 1 FROM projects WHERE id=$1 AND company_id=$2', project_id, company):
            raise AdvanceError('Проект недоступен.')
        if employee_id is not None and not await connection.fetchval(
                'SELECT 1 FROM company_members WHERE user_id=$1 AND company_id=$2', employee_id, company):
            raise AdvanceError('Сотрудник недоступен.')

    @staticmethod
    def _conditions(company, *, q='', project_id=None, employee_id=None,
                    duplicate_status=None, payment_source=None, start=None, end=None,
                    document_type=None):
        if not isinstance(q, str) or len(q) > 200:
            raise AdvanceError('Поисковый запрос должен быть не длиннее 200 символов.')
        clauses, values = ['d.company_id=$1'], [company]

        def add(sql, value):
            values.append(value)
            clauses.append(sql.replace('?', f'${len(values)}'))

        if q.strip():
            add("(d.vendor ILIKE ? OR d.external_document_number ILIKE ? OR d.incoming_number ILIKE ?)", '%' + q.strip() + '%')
        if project_id is not None:
            add('d.project_id=?', project_id)
        if employee_id is not None:
            add('d.uploaded_by_user_id=?', employee_id)
        if duplicate_status is not None:
            if duplicate_status not in ('none', 'exact', 'probable', 'not_checked', 'flagged'):
                raise AdvanceError('Неизвестный статус дубля.')
            if duplicate_status == 'flagged':
                clauses.append("d.duplicate_status IN ('exact','probable')")
            else:
                add('d.duplicate_status=?', duplicate_status)
        if payment_source is not None:
            if payment_source == 'advance':
                clauses.append('ae.id IS NOT NULL')
            elif payment_source == 'card':
                clauses.append('ce.id IS NOT NULL')
            elif payment_source in ('unassigned', 'none'):
                clauses.append('ae.id IS NULL AND ce.id IS NULL')
            else:
                raise AdvanceError('Неизвестный источник оплаты.')
        if start is not None:
            if not isinstance(start, datetime) or start.tzinfo is None:
                raise AdvanceError('Начало периода должно содержать часовой пояс.')
            add('d.created_at>=?', start)
        if end is not None:
            if not isinstance(end, datetime) or end.tzinfo is None:
                raise AdvanceError('Конец периода должен содержать часовой пояс.')
            add('d.created_at<?', end)
        if start is not None and end is not None:
            _period(start, end)
        if document_type is not None:
            add('d.document_type=?', document_type)
        return ' AND '.join(clauses), values

    async def documents(self, actor, *, page=1, per_page=25, q='', project_id=None,
                        employee_id=None, duplicate_status=None, payment_source=None,
                        start=None, end=None, document_type=None, sort='newest'):
        _page(page, per_page)
        orders = {
            'newest': 'd.created_at DESC,d.id DESC',
            'oldest': 'd.created_at ASC,d.id ASC',
            'amount_desc': 'd.total_amount DESC NULLS LAST,d.id DESC',
            'amount_asc': 'd.total_amount ASC NULLS LAST,d.id ASC',
            'vendor': 'd.vendor ASC NULLS LAST,d.id ASC',
            'name': 'd.vendor ASC NULLS LAST,d.id ASC',
        }
        if sort not in orders:
            raise AdvanceError('Неизвестный порядок сортировки документов.')
        async with get_pool().acquire() as c, c.transaction(isolation='repeatable_read', readonly=True):
            company = await self._manager(c, actor)
            await self._scope(c, company, project_id, employee_id)
            where, params = self._conditions(
                company, q=q, project_id=project_id, employee_id=employee_id,
                duplicate_status=duplicate_status, payment_source=payment_source,
                start=start, end=end, document_type=document_type,
            )
            query = _DOCUMENT_SELECT + ' WHERE ' + where
            total = await c.fetchval('SELECT count(*) FROM (' + query + ') selected', *params)
            rows = await c.fetch(query + f' ORDER BY {orders[sort]} LIMIT ${len(params)+1} OFFSET ${len(params)+2}',
                                 *params, per_page, (page - 1) * per_page)
            return {'items': [dict(row) for row in rows], 'total': total, 'page': page, 'per_page': per_page}

    async def document(self, actor, document_id):
        async with get_pool().acquire() as c, c.transaction(isolation='repeatable_read', readonly=True):
            company = await self._manager(c, actor)
            row = await c.fetchrow(_DOCUMENT_SELECT + ' WHERE d.company_id=$1 AND d.id=$2', company, document_id)
            if row is None:
                raise AdvanceError('Документ недоступен.')
            result = dict(row)
            result['items'] = [dict(item) for item in await c.fetch(
                'SELECT id,document_id,line_no,name,quantity,price,line_total FROM document_items WHERE document_id=$1 ORDER BY line_no,id', document_id)]
            result['funding_history'] = [dict(item) for item in await c.fetch("""
                SELECT 'advance' AS source,e.id,e.advance_id AS source_id,e.amount,e.created_at,
                       e.approved_at,e.rejected_at,e.rejection_reason FROM advance_expenses e WHERE e.document_id=$1
                UNION ALL SELECT 'card',e.id,e.card_id,e.amount,e.created_at,e.approved_at,e.rejected_at,e.rejection_reason
                FROM card_expenses e WHERE e.document_id=$1 ORDER BY created_at,id
            """, document_id)]
            return result

    async def duplicate_keep(self, actor, document_id):
        async with get_pool().acquire() as c, c.transaction():
            company = await self._manager(c, actor, lock=True)
            row = await c.fetchrow('SELECT id,duplicate_status FROM documents WHERE id=$1 AND company_id=$2 FOR UPDATE', document_id, company)
            if row is None:
                raise AdvanceError('Документ недоступен.')
            if row['duplicate_status'] == 'none':
                return {'id': document_id, 'duplicate_status': 'none'}
            await c.execute("UPDATE documents SET duplicate_status='none',duplicate_of_document_id=NULL,duplicate_checked_at=clock_timestamp() WHERE id=$1", document_id)
            await self._audit(c, actor, company, 'document.duplicate_keep', document_id)
            return {'id': document_id, 'duplicate_status': 'none'}

    async def duplicate_delete(self, actor, document_id, source=False):
        async with get_pool().acquire() as c, c.transaction():
            company = await self._manager(c, actor, lock=True)
            original = await c.fetchrow('SELECT id,duplicate_status,duplicate_of_document_id FROM documents WHERE id=$1 AND company_id=$2', document_id, company)
            if original is None or original['duplicate_status'] not in ('exact', 'probable'):
                raise AdvanceError('Дубликат недоступен или уже обработан.')
            target = original['duplicate_of_document_id'] if source else document_id
            if target is None:
                raise AdvanceError('Исходная запись для дубля не найдена.')
            locked = await c.fetch('SELECT id,duplicate_status,duplicate_of_document_id FROM documents WHERE id=ANY($1::bigint[]) AND company_id=$2 ORDER BY id FOR UPDATE', sorted({document_id, target}), company)
            rows = {row['id']: row for row in locked}
            current = rows.get(document_id)
            if (target not in rows or current is None
                    or current['duplicate_status'] not in ('exact', 'probable')
                    or (source and current['duplicate_of_document_id'] != target)):
                raise AdvanceError('Состояние дубля изменилось. Обновите список.')
            if await c.fetchval('SELECT EXISTS(SELECT 1 FROM advance_expenses WHERE document_id=$1) OR EXISTS(SELECT 1 FROM card_expenses WHERE document_id=$1)', target):
                raise AdvanceError('Документ связан с финансовой историей и не может быть удалён.')
            await c.execute("UPDATE documents SET duplicate_status='none',duplicate_of_document_id=NULL,duplicate_checked_at=clock_timestamp() WHERE company_id=$1 AND duplicate_of_document_id=$2", company, target)
            await c.execute('DELETE FROM documents WHERE id=$1 AND company_id=$2', target, company)
            await self._audit(c, actor, company, 'document.duplicate_delete', target)
            return {'id': target, 'deleted': True}

    @staticmethod
    async def _audit(c, actor, company, action, document_id):
        # Table and shared event schema are created by the workspace migration.
        await c.execute("INSERT INTO workspace_events(actor_id,company_id,action,entity_type,entity_id,details) VALUES($1,$2,$3,'document',$4,'{}'::jsonb)", actor, company, action, document_id)

    @staticmethod
    async def _money(c, actor, start, end, *, project_id=None, employee_id=None):
        scope = {'project_id': project_id, 'employee_id': employee_id, 'connection': c}
        cash = await AdvanceService().report(actor, start, end, **scope)
        cash['card_rows'] = [dict(row) for row in await BusinessCardService().report(actor, start, end, **scope)]
        cash['card_spent'] = sum((r['amount'] for r in cash['card_rows'] if r['accepted']), Decimal(0))
        cash['card_pending'] = sum((r['amount'] for r in cash['card_rows'] if not r['accepted']), Decimal(0))
        cash['cash_on_hand'], cash['employee_debt'] = cash_and_employee_debt(cash)
        return cash

    async def overview(self, actor, start, end):
        """Dashboard aggregates avoid loading document bodies or line items."""
        _period(start, end)
        async with get_pool().acquire() as c, c.transaction(isolation='repeatable_read', readonly=True):
            company = await self._manager(c, actor)
            currencies = await c.fetch("""
                SELECT currency,count(*) AS document_count,COALESCE(sum(total_amount),0) AS total_amount,
                    count(*) FILTER (WHERE duplicate_status='exact') AS exact_duplicates,
                    count(*) FILTER (WHERE duplicate_status='probable') AS probable_duplicates
                FROM documents WHERE company_id=$1 AND created_at>=$2 AND created_at<$3
                GROUP BY currency
            """, company, start, end)
            # Unresolved duplicate warnings must survive period changes and project archival.
            duplicate_review = dict(await c.fetchrow("""
                SELECT count(*) AS total,
                    count(*) FILTER (WHERE d.duplicate_status='exact') AS exact,
                    count(*) FILTER (WHERE d.duplicate_status='probable') AS probable,
                    count(*) FILTER (WHERE p.status='archived' OR source_p.status='archived') AS involving_archived_projects
                FROM documents d JOIN projects p ON p.id=d.project_id AND p.company_id=d.company_id
                LEFT JOIN documents source ON source.id=d.duplicate_of_document_id AND source.company_id=d.company_id
                LEFT JOIN projects source_p ON source_p.id=source.project_id AND source_p.company_id=d.company_id
                WHERE d.company_id=$1 AND d.duplicate_status IN ('exact','probable')
            """, company))
            money = await self._money(c, actor, start, end)
            projects = await c.fetch("SELECT id,name FROM projects WHERE company_id=$1 AND status='active' ORDER BY lower(name),id", company)
            project_documents = await c.fetch('''SELECT d.project_id,d.currency,count(*) AS document_count,
                COALESCE(sum(d.total_amount),0) AS total_amount
                FROM documents d JOIN projects p ON p.id=d.project_id AND p.company_id=d.company_id
                WHERE d.company_id=$1 AND p.status='active' AND d.created_at>=$2 AND d.created_at<$3
                GROUP BY d.project_id,d.currency''', company, start, end)
            employee_count = await c.fetchval("SELECT count(*) FROM company_members WHERE company_id=$1 AND role='employee' AND status='active'", company)
            generated_at = await c.fetchval('SELECT clock_timestamp()')
        open_projects = []
        cash_by_project, cards_by_project, docs_by_project = defaultdict(list), defaultdict(list), defaultdict(list)
        for row in money['rows']:
            cash_by_project[row['project_id']].append(row)
        for row in money['card_rows']:
            cards_by_project[row['project_id']].append(row)
        for row in project_documents:
            docs_by_project[row['project_id']].append(row)
        for project in projects:
            cash_rows = cash_by_project[project['id']]
            card_rows = cards_by_project[project['id']]
            docs = docs_by_project[project['id']]
            row = {'project_id': project['id'], 'project_name': project['name'],
                   'document_count': sum(d['document_count'] for d in docs),
                   'document_total': sum((d['total_amount'] for d in docs if d['currency'] == 'RUB'), Decimal(0)),
                   'totals_by_currency': {d['currency']: d['total_amount'] for d in docs},
                   **{key: sum((r[key] for r in cash_rows), Decimal(0)) for key in CASH_FIELDS}}
            row['cash'], row['debt'] = cash_and_employee_debt({'rows': cash_rows})
            row['card_spent'] = sum((r['amount'] for r in card_rows if r['accepted']), Decimal(0))
            row['card_pending'] = sum((r['amount'] for r in card_rows if not r['accepted']), Decimal(0))
            row['review_total'] = row['pending_expenses'] + row['pending_returns'] + row['card_pending']
            open_projects.append(row)
        currency_totals = {row['currency']: row['total_amount'] for row in currencies}
        rub_total = currency_totals.get('RUB', Decimal(0))
        count = sum(row['document_count'] for row in currencies)
        return {
            'start': start, 'end': end, 'start_at': start, 'end_at': end, 'generated_at': generated_at,
            'document_count': count, 'documents': count, 'document_total': rub_total,
            'project_count': len(projects), 'employee_count': employee_count,
            'open_projects': open_projects,
            'total_amount': rub_total, 'currency': 'RUB', 'totals_by_currency': currency_totals,
            'exact_duplicates': sum(row['exact_duplicates'] for row in currencies),
            'probable_duplicates': sum(row['probable_duplicates'] for row in currencies),
            'duplicate_review': duplicate_review,
            **money['totals'], 'card_spent': money['card_spent'], 'card_pending': money['card_pending'],
            'cash': money['cash_on_hand'], 'debt': money['employee_debt'], 'money': money,
        }

    async def report(self, actor, start, end, *, project_id=None, employee_id=None):
        _period(start, end)
        async with get_pool().acquire() as c, c.transaction(isolation='repeatable_read', readonly=True):
            company = await self._manager(c, actor)
            await self._scope(c, company, project_id, employee_id)
            where, params = self._conditions(company, start=start, end=end,
                                              project_id=project_id, employee_id=employee_id)
            documents = [dict(row) for row in await c.fetch(_DOCUMENT_SELECT + ' WHERE ' + where + ' ORDER BY d.id', *params)]
            # No display limit is applied to reports or exports.
            items = [dict(row) for row in await c.fetch('SELECT di.* FROM document_items di JOIN documents d ON d.id=di.document_id WHERE ' + where + ' ORDER BY di.document_id,di.line_no,di.id', *params)]
            cash = await self._money(c, actor, start, end, project_id=project_id, employee_id=employee_id)
            generated_at = await c.fetchval('SELECT clock_timestamp()')

        currency_totals = defaultdict(Decimal)
        for doc in documents:
            currency_totals[doc['currency']] += doc['total_amount'] or Decimal(0)
        summary = {
            'start_at': start, 'end_at': end, 'document_count': len(documents), 'documents': len(documents),
            'total_amount': currency_totals.get('RUB', Decimal(0)), 'currency': 'RUB',
            'totals_by_currency': dict(currency_totals),
            'exact_duplicates': sum(d['duplicate_status'] == 'exact' for d in documents),
            'probable_duplicates': sum(d['duplicate_status'] == 'probable' for d in documents),
        }
        projects = self._group(documents, 'project_id')
        employees = self._group(documents, 'uploaded_by_user_id')
        return {
            'start': start, 'end': end, 'generated_at': generated_at,
            'filters': {'project_id': project_id, 'employee_id': employee_id},
            'summary': summary, 'overview': summary, 'projects': projects, 'employees': employees,
            'documents': documents, 'items': items,
            'duplicates': [d for d in documents if d['duplicate_status'] in ('exact', 'probable')],
            'money': cash,
        }

    @staticmethod
    def _group(documents, field):
        groups = {}
        for d in documents:
            key = d[field]
            if key not in groups:
                groups[key] = {
                    'id': key, 'project_id': d['project_id'] if field == 'project_id' else None,
                    'project_name': d['project_name'] if field == 'project_id' else None,
                    'user_id': d['uploaded_by_user_id'] if field == 'uploaded_by_user_id' else None,
                    'employee_name': d['uploaded_by_name'] if field == 'uploaded_by_user_id' else None,
                    'username': d['uploaded_by_username'] if field == 'uploaded_by_user_id' else None,
                    'document_count': 0, 'total_amount': Decimal(0), 'currency': 'RUB',
                    'totals_by_currency': defaultdict(Decimal),
                    'exact_duplicate_count': 0, 'probable_duplicate_count': 0,
                }
            row = groups[key]
            row['document_count'] += 1
            row['totals_by_currency'][d['currency']] += d['total_amount'] or Decimal(0)
            if d['currency'] == 'RUB':
                row['total_amount'] += d['total_amount'] or Decimal(0)
            row['exact_duplicate_count'] += d['duplicate_status'] == 'exact'
            row['probable_duplicate_count'] += d['duplicate_status'] == 'probable'
        for row in groups.values():
            row['totals_by_currency'] = dict(row['totals_by_currency'])
        return list(groups.values())

    async def export(self, actor, start, end, *, project_id=None, employee_id=None):
        result = await self.report(actor, start, end, project_id=project_id, employee_id=employee_id)
        summary = SimpleNamespace(**result['summary'])
        as_objects = lambda values: [SimpleNamespace(**row) for row in values]
        _, payload = build_manager_report_workbook(
            f'{start:%Y-%m-%d} — {end:%Y-%m-%d} UTC (конец не включён)', summary,
            as_objects(result['projects']), as_objects(result['employees']),
            as_objects(result['duplicates']), as_objects(result['documents']),
            as_objects(result['items']), money_report=result['money'],
        )
        workbook = load_workbook(BytesIO(payload))
        overview = workbook['Overview']
        _append_row(overview, ['До (не включая), UTC', end.isoformat()])
        _append_row(overview, ['Сформирован, UTC', result['generated_at'].isoformat()])
        _append_row(overview, ['Проект ID', project_id or 'Все'])
        _append_row(overview, ['Сотрудник ID', employee_id or 'Все'])
        _append_row(overview, ['Основание периода документов', 'Дата загрузки, UTC'])
        _append_row(overview, ['Основание периода денег', 'Дата подтверждения, UTC'])
        overview['A4'] = 'Сумма RUB'
        workbook['Projects']['D1'] = 'Сумма RUB'
        workbook['Employees']['E1'] = 'Сумма RUB'
        for currency, total in result['summary']['totals_by_currency'].items():
            _append_row(overview, [f'Документы: {currency}', float(total)])
        for sheet_name, rows, column in (
                ('Documents', result['documents'], 11), ('Duplicates', result['duplicates'], 11)):
            workbook[sheet_name].cell(1, column, 'Валюта')
            for row_number, doc in enumerate(rows, 2):
                cell = workbook[sheet_name].cell(row_number, column, _ILLEGAL_XML.sub('\ufffd', doc['currency']))
                cell.data_type = 's'
        output = BytesIO()
        workbook.save(output)
        return f'fiinops-report-{start:%Y%m%d}-{end:%Y%m%d}.xlsx', output.getvalue()
