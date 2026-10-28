"""Integration tests create/drop a unique database, never mutate POSTGRES_DB."""
import asyncio
import os
import unittest
from contextlib import asynccontextmanager
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import AsyncMock, patch
from uuid import uuid4

import asyncpg
from aiogram.types import User

from app.config import get_settings
from app.handlers import documents as handlers
from app.handlers.start import start_command
from app.schemas.document import DocumentSchema
from app.services.companies import CompanyAccessError, CompanyService
from app.services.database import init_db, close_db
from app.services.documents import DocumentService
from app.services.projects import ProjectService
from app.state.pending_actions import set_pending_action, pop_pending_action
from app.state.pending_documents import (
    PendingDocument, begin_document_flow, get_pending_document,
    clear_document_flow, store_pending_document, has_active_document_flow,
)
from app.ui.projects import PROJECT_CALLBACK_PREFIX, DOCUMENT_DUPLICATE_SAVE_CALLBACK, PROJECT_CANCEL_CALLBACK


from db_fixture import DatabaseFixture


@unittest.skipUnless(os.environ.get('RUN_DATABASE_TESTS') == '1', 'Set RUN_DATABASE_TESTS=1 and PostgreSQL connection variables')
class DatabaseTests(DatabaseFixture):






    async def test_single_invite_allows_only_one_concurrent_join(self):
        code = await self.service.create_invite(self.manager, 'employee')
        users = [self.user(1010+i) for i in range(5)]
        results = await asyncio.wait_for(asyncio.gather(*(self.service.join_company(user, code) for user in users), return_exceptions=True), 15)
        self.assertEqual(sum(not isinstance(r, Exception) for r in results), 1)
        self.assertEqual(sum(isinstance(r, CompanyAccessError) for r in results), 4)
        async with self.pool.acquire() as c:
            self.assertEqual(await c.fetchval("SELECT count(*) FROM company_members WHERE role='employee'"), 2)

    async def test_employee_limit_serializes_owner_assignments_and_join(self):
        for user_id in range(1020, 1028):
            user = self.user(user_id)
            uid = await self.service.ensure_platform_user(user)
            await self.service.assign_user_to_company_by_owner(self.owner.id, uid, self.company.id, 'employee')
        code = await self.service.create_invite(self.manager, 'employee')
        uid = await self.service.ensure_platform_user(self.user(1030))
        results = await asyncio.wait_for(asyncio.gather(
            self.service.join_company(self.user(1031), code),
            self.service.assign_user_to_company_by_owner(self.owner.id, uid, self.company.id, 'employee'),
            return_exceptions=True,
        ), 15)
        self.assertEqual(sum(not isinstance(r, Exception) for r in results), 1)
        async with self.pool.acquire() as c:
            self.assertEqual(await c.fetchval("SELECT count(*) FROM company_members WHERE role='employee' AND status='active'"), 10)

    async def test_expired_invite_rejected(self):
        code = await self.service.create_invite(self.manager, 'employee')
        async with self.pool.acquire() as c:
            await c.execute("UPDATE company_invites SET expires_at=NOW()-interval '1 second' WHERE code=$1", code)
        with self.assertRaises(CompanyAccessError):
            await self.service.join_company(self.user(1040), code)

    async def test_invite_code_collision_retries_inside_savepoint(self):
        code = await self.service.create_invite(self.manager, 'employee')
        with patch('app.services.companies._generate_invite_code', side_effect=[code, 'FRESHCODE']):
            self.assertEqual(await self.service.create_invite(self.manager, 'employee'), 'FRESHCODE')

    async def test_owner_rotation_revokes_cached_privilege_without_old_user_visit(self):
        get_settings().bot_owner_telegram_id = 1999
        self.assertFalse(await self.service.is_platform_owner(self.owner.id))
        self.assertEqual((await self.service.get_user_context(self.owner.id)).menu_kind, 'employee')
        with self.assertRaises(CompanyAccessError):
            await self.service.create_company(self.owner, 'Denied')
        self.assertFalse(await self.service.is_platform_owner(1999))
        await self.service.ensure_platform_user(self.user(1999))
        self.assertTrue(await self.service.is_platform_owner(1999))
        await self.service.ensure_platform_user(self.owner)
        async with self.pool.acquire() as c:
            self.assertEqual(await c.fetchval('SELECT system_role FROM users WHERE telegram_id=$1', self.owner.id), 'user')

    async def test_owner_cannot_join_company_and_manager_cannot_demote_via_invite(self):
        code = await self.service.create_invite(self.manager, 'employee')
        for user in (self.owner, self.manager):
            with self.assertRaises(CompanyAccessError):
                await self.service.join_company(user, code)

    async def test_draft_claim_is_atomic(self):
        claims = await asyncio.gather(*(begin_document_flow(self.employee.id) for _ in range(4)))
        self.assertEqual(sum(c is not None for c in claims), 1)

    async def test_draft_and_action_survive_pool_restart_with_typed_duplicate_fields(self):
        draft = await self.make_draft()
        draft.duplicate_check = await DocumentService().find_company_duplicate_document(self.employee, self.project, self.document)
        draft.selected_project_id = self.project.id
        await store_pending_document(self.employee.id, draft)
        await set_pending_action(self.employee.id, 'rename_project', {'project_id': self.project.id})
        await close_db()
        self.pool = await init_db()
        restored = await get_pending_document(self.employee.id)
        self.assertEqual(restored.upload_id, draft.upload_id)
        self.assertEqual(restored.extracted_document, draft.extracted_document)
        self.assertEqual(restored.duplicate_check, draft.duplicate_check)
        action = await pop_pending_action(self.employee.id)
        self.assertEqual(action.payload['project_id'], self.project.id)
        self.assertIsNone(await pop_pending_action(self.employee.id))

    async def test_expired_state_is_unusable_and_cannot_overwrite_new_draft(self):
        draft = await self.make_draft()
        await set_pending_action(self.employee.id, 'create_project')
        async with self.pool.acquire() as c:
            await c.execute("UPDATE pending_document_flows SET expires_at=NOW()-interval '1 second'")
            await c.execute("UPDATE pending_actions SET expires_at=NOW()-interval '1 second'")
        self.assertIsNone(await get_pending_document(self.employee.id))
        self.assertFalse(await has_active_document_flow(self.employee.id))
        self.assertIsNone(await pop_pending_action(self.employee.id))
        new_key = await begin_document_flow(self.employee.id)
        self.assertNotEqual(new_key, draft.upload_id)
        self.assertFalse(await store_pending_document(self.employee.id, draft))
        await clear_document_flow(self.employee.id, draft.upload_id)
        self.assertTrue(await has_active_document_flow(self.employee.id))

    async def test_action_consumed_once_by_concurrent_messages(self):
        await set_pending_action(self.employee.id, 'create_project')
        values = await asyncio.gather(*(pop_pending_action(self.employee.id) for _ in range(3)))
        self.assertEqual(sum(v is not None for v in values), 1)

    async def test_idempotency_key_prevents_double_insert_but_allows_intentional_duplicate(self):
        service = DocumentService()
        check = await service.find_company_duplicate_document(self.employee, self.project, self.document)
        async def save(key):
            return await service.save_document(self.employee, self.project, 'preview', self.document, check, upload_id=key)
        ids = await asyncio.gather(save('same-upload'), save('same-upload'))
        self.assertEqual(ids[0], ids[1])
        self.assertEqual(await self.document_count(), 1)
        async with self.pool.acquire() as c:
            self.assertEqual(await c.fetchval('SELECT count(*) FROM document_items'), 1)
        self.assertNotEqual(await save('intentional-new-upload'), ids[0])
        self.assertEqual(await self.document_count(), 2)

    async def test_concurrent_project_clicks_save_once(self):
        draft = await self.make_draft()
        callbacks = [self.callback(f'{PROJECT_CALLBACK_PREFIX}{self.project.id}:{draft.upload_id}') for _ in range(3)]
        await asyncio.wait_for(asyncio.gather(*(handlers.process_project_selection(c) for c in callbacks)), 15)
        self.assertEqual(await self.document_count(), 1)
        self.assertIsNone(await get_pending_document(self.employee.id))

    async def test_duplicate_confirmation_survives_restart_and_saves_once(self):
        await DocumentService().save_document(self.employee, self.project, 'preview', self.document, upload_id='original')
        draft = await self.make_draft()
        await handlers.process_project_selection(self.callback(f'{PROJECT_CALLBACK_PREFIX}{self.project.id}:{draft.upload_id}'))
        self.assertEqual(await self.document_count(), 1)
        await close_db()
        self.pool = await init_db()
        callbacks = [self.callback(f'{DOCUMENT_DUPLICATE_SAVE_CALLBACK}:{draft.upload_id}') for _ in range(2)]
        await asyncio.gather(*(handlers.duplicate_save_callback(cb) for cb in callbacks))
        self.assertEqual(await self.document_count(), 2)

    async def test_old_buttons_do_not_save_or_cancel_new_upload(self):
        draft = await self.make_draft()
        stale = 'previous-upload'
        await handlers.process_project_selection(self.callback(f'{PROJECT_CALLBACK_PREFIX}{self.project.id}:{stale}'))
        await handlers.cancel_project_selection(self.callback(f'{PROJECT_CANCEL_CALLBACK}:{stale}'))
        self.assertEqual(await self.document_count(), 0)
        self.assertEqual((await get_pending_document(self.employee.id)).upload_id, draft.upload_id)

    async def test_cross_company_project_access_is_denied(self):
        company = await self.service.create_company(self.owner, 'Other company')
        async with self.pool.acquire() as c:
            pid = await c.fetchval('INSERT INTO projects(company_id,name,created_by_user_id) SELECT $1,$2,id FROM users WHERE telegram_id=$3 RETURNING id', company.id, 'Other project', self.owner.id)
        self.assertIsNone(await ProjectService().get_active_project(self.employee.id, pid))
        from app.services.projects import Project
        with self.assertRaises(CompanyAccessError):
            await DocumentService().save_document(self.employee, Project(pid, company.id, 'Other project', 'active'), 'preview', self.document, upload_id='denied')

    async def test_start_returns_owner_and_employee_menus(self):
        for user in (self.owner, self.employee):
            message = SimpleNamespace(from_user=user, answer=AsyncMock())
            await start_command(message)
            self.assertEqual(message.answer.await_count, 1)
            self.assertIsNotNone(message.answer.call_args.kwargs['reply_markup'])

    async def test_migration_is_repeatable_and_keeps_existing_documents(self):
        await DocumentService().save_document(self.employee, self.project, 'preview', self.document, upload_id='migration-test')
        await close_db()
        self.pool = await init_db()
        self.assertEqual(await self.document_count(), 1)
        async with self.pool.acquire() as c:
            self.assertEqual(await c.fetchval('SELECT count(*) FROM schema_migrations'), len(list((Path(__file__).resolve().parent.parent / 'db/migrations').glob('*.sql'))))

    async def test_photo_cleanup_after_success_and_ocr_failure(self):
        from app.services.ocr_space import OCRSpaceError
        from app.services import document_processing as processing
        import tempfile
        from uuid import uuid4
        @asynccontextmanager
        async def no_typing(**kwargs):
            yield
        async def download(file, destination):
            destination.write(b'raw-image')
        for failure in (False, True):
            with tempfile.TemporaryDirectory() as folder:
                path = Path(folder) / 'prepared.jpg'
                path.write_bytes(b'fake prepared image')
                message = SimpleNamespace(message_id=uuid4().hex, from_user=self.employee, photo=[SimpleNamespace(file_id='x')], chat=SimpleNamespace(id=1), bot=SimpleNamespace(download=download), answer=AsyncMock())
                ocr = SimpleNamespace(extract_text=AsyncMock(side_effect=OCRSpaceError('failed')) if failure else AsyncMock(return_value='Кассовый чек'))
                with patch('app.services.document_flow.get_provider_for_user', AsyncMock(return_value='ocr_space')), patch.object(handlers.ChatActionSender, 'typing', no_typing), patch.object(processing, 'prepare_image', return_value=path), patch.object(processing, 'create_recognizer', return_value=ocr), patch.object(processing.DeepSeekService, 'extract_document', AsyncMock(return_value=self.document.model_dump())):
                    await handlers.process_photo(message)
                self.assertFalse(path.exists())
                self.assertEqual(await has_active_document_flow(self.employee.id), not failure)
                await clear_document_flow(self.employee.id)

    async def test_pdf_preview_project_save_and_failure_cleanup(self):
        from app.services.pdf_files import PDFError
        from app.services import document_processing as processing
        @asynccontextmanager
        async def no_typing(**kwargs):
            yield
        async def download(file, destination):
            destination.write(b'%PDF-test')
        message = SimpleNamespace(message_id=1, from_user=self.employee, photo=None,
            document=SimpleNamespace(mime_type='application/pdf', file_name='receipt.pdf'),
            caption='все', chat=SimpleNamespace(id=1), bot=SimpleNamespace(download=download), answer=AsyncMock())
        with patch('app.services.document_flow.get_provider_for_user', AsyncMock(return_value='ocr_space')), patch.object(handlers.ChatActionSender, 'typing', no_typing), patch.object(processing, 'extract_pdf_bytes', AsyncMock(return_value='Кассовый чек')), patch.object(processing.DeepSeekService, 'extract_document', AsyncMock(return_value=self.document.model_dump())):
            await handlers.process_pdf(message)
        draft = await get_pending_document(self.employee.id)
        self.assertEqual(draft.source_type, 'pdf')
        self.assertEqual(await self.document_count(), 0)
        cb = self.callback(f'{PROJECT_CALLBACK_PREFIX}{self.project.id}:{draft.upload_id}')
        await handlers.process_project_selection(cb)
        self.assertEqual(await self.document_count(), 1)
        self.assertEqual(await self.pool.fetchval('SELECT source_type FROM documents'), 'pdf')
        message.message_id = 2
        with patch('app.services.document_flow.get_provider_for_user', AsyncMock(return_value='ocr_space')), patch.object(handlers.ChatActionSender, 'typing', no_typing), patch.object(processing, 'extract_pdf_bytes', AsyncMock(side_effect=PDFError('bad PDF'))):
            await handlers.process_pdf(message)
        self.assertFalse(await has_active_document_flow(self.employee.id))
        self.assertEqual(await self.document_count(), 1)

    async def test_save_failure_keeps_draft_for_idempotent_retry(self):
        draft = await self.make_draft()
        cb = self.callback(f'{PROJECT_CALLBACK_PREFIX}{self.project.id}:{draft.upload_id}')
        with patch.object(handlers.document_service, 'save_for_user', AsyncMock(side_effect=ConnectionError('temporary'))):
            with self.assertLogs(handlers.logger, level='ERROR'):
                await handlers.process_project_selection(cb)
        self.assertIsNotNone(await get_pending_document(self.employee.id))
        self.assertEqual(await self.document_count(), 0)
        await handlers.process_project_selection(cb)
        self.assertEqual(await self.document_count(), 1)

    async def test_handler_dispatch_routes_document_callbacks_with_injected_data(self):
        # Exercise aiogram routing/signature inspection, not only direct function calls.
        from aiogram import Bot, Dispatcher
        from aiogram.types import Update
        from app.handlers.documents import router
        # Use a fresh observer via a fresh Router to avoid attaching a shared router
        # to multiple dispatchers in the process.
        from aiogram import Router
        isolated = Router()
        for handler in router.callback_query.handlers:
            isolated.callback_query.register(handler.callback, *[f.callback for f in handler.filters])
        dispatcher = Dispatcher()
        dispatcher.include_router(isolated)
        draft = await self.make_draft()
        bot = Bot('123456:TEST_ONLY')
        update = Update.model_validate({'update_id': 1, 'callback_query': {'id': 'test', 'from': {'id': self.employee.id, 'is_bot': False, 'first_name': 'Employee'}, 'chat_instance': 'test', 'data': f'{PROJECT_CALLBACK_PREFIX}{self.project.id}:{draft.upload_id}', 'message': {'message_id': 1, 'date': 0, 'chat': {'id': self.employee.id, 'type': 'private'}, 'text': 'Choose'}}})
        with patch.object(bot.session, 'make_request', AsyncMock(return_value=True)):
            await dispatcher.feed_update(bot, update)
        self.assertEqual(await self.document_count(), 1)
        await bot.session.close()

    async def test_upload_id_cannot_be_reused_for_a_different_project(self):
        from app.services.documents import DocumentValidationError
        service = DocumentService()
        await service.save_document(self.employee, self.project, 'preview', self.document, upload_id='one-project')
        other = await ProjectService().create_project(self.manager.id, 'Different project')
        with self.assertRaises(DocumentValidationError):
            await service.save_document(self.employee, other, 'preview', self.document, upload_id='one-project')
        self.assertEqual(await self.document_count(), 1)


    async def test_advance_lifecycle_posts_costs_immediately_and_confirms_returns(self):
        from app.services.advances import AdvanceService
        s=AdvanceService();manager,employee=await self.finance_ids()
        aid=await s.issue(manager,employee,self.project.id,'150','Materials','issue-1')
        self.assertEqual(await s.issue(manager,employee,self.project.id,'150','Materials','issue-1'),aid)
        await s.receive(employee,aid);await s.receive(employee,aid)
        document=await DocumentService().save_document(self.employee,self.project,'preview',self.document,upload_id='finance-check')
        eid=await s.attach(employee,aid,document)
        self.assertEqual((await s.card(employee,aid))['balance'],50)
        self.assertEqual(await s.attach(employee,aid,document),eid)
        await s.approve_expense(manager,aid,eid);await s.approve_expense(manager,aid,eid)
        self.assertEqual((await s.card(employee,aid))['balance'],50)
        rid=await s.request_return(employee,aid,50,'return-1')
        self.assertEqual((await s.card(employee,aid))['balance'],50)
        await s.approve_return(manager,aid,rid);await s.approve_return(manager,aid,rid)
        self.assertEqual((await s.card(employee,aid))['balance'],0)
        async with self.pool.acquire() as c:
            self.assertEqual(await c.fetchval('SELECT count(*) FROM advance_events WHERE advance_id=$1',aid),5)

    async def test_advance_permissions_and_unreceived_expense(self):
        from app.services.advances import AdvanceService,AdvanceError
        s=AdvanceService();manager,employee=await self.finance_ids()
        with self.assertRaises(AdvanceError): await s.issue(employee,employee,self.project.id,100,'x','denied')
        aid=await s.issue(manager,employee,self.project.id,100,'x','valid')
        with self.assertRaises(AdvanceError): await s.receive(manager,aid)
        with self.assertRaises(AdvanceError): await s.attach(employee,aid,123)
        with self.assertRaises(AdvanceError): await s.approve_return(employee,aid,1)
        outsider=await self.service.ensure_platform_user(self.user(1888))
        with self.assertRaises(AdvanceError): await s.card(outsider,aid)

    async def test_advance_double_approval_and_return_limit(self):
        from app.services.advances import AdvanceService,AdvanceError
        s=AdvanceService();manager,employee=await self.finance_ids()
        aid=await s.issue(manager,employee,self.project.id,100,'x','valid')
        await s.receive(employee,aid)
        one=await s.request_return(employee,aid,80,'one')
        two=await s.request_return(employee,aid,80,'two')
        results=await asyncio.gather(s.approve_return(manager,aid,one),s.approve_return(manager,aid,two),return_exceptions=True)
        self.assertEqual(sum(isinstance(r,AdvanceError) for r in results),1)
        self.assertEqual((await s.card(employee,aid))['balance'],20)

    async def test_expense_cannot_cover_two_advances_and_overspend_is_visible(self):
        from app.services.advances import AdvanceService,AdvanceError
        s=AdvanceService();manager,employee=await self.finance_ids()
        one=await s.issue(manager,employee,self.project.id,50,'x','one')
        two=await s.issue(manager,employee,self.project.id,50,'x','two')
        await s.receive(employee,one);await s.receive(employee,two)
        doc=await DocumentService().save_document(self.employee,self.project,'preview',self.document,upload_id='check')
        eid=await s.attach(employee,one,doc)
        self.assertEqual(await s.attach(employee,one,doc),eid)
        with self.assertRaises(AdvanceError): await s.attach(employee,two,doc)
        await s.approve_expense(manager,one,eid)
        self.assertEqual((await s.card(employee,one))['balance'],-50)

    async def test_advance_rejects_invalid_money_and_conflicting_retry(self):
        from app.services.advances import AdvanceService,AdvanceError,money
        for value in ('NaN','Infinity','-1','0','0.001','1000000000000','bad'):
            with self.assertRaises(AdvanceError): money(value)
        s=AdvanceService();manager,employee=await self.finance_ids()
        await s.issue(manager,employee,self.project.id,'10,50','x','key')
        with self.assertRaises(AdvanceError): await s.issue(manager,employee,self.project.id,11,'x','key')

    async def test_finance_telegram_flow_selects_employee_project_and_amount(self):
        from app.handlers import advances as ui
        manager,employee=await self.finance_ids()
        def cb(user,data):
            return SimpleNamespace(id='unique-ui-request',from_user=user,data=data,answer=AsyncMock(),message=SimpleNamespace(answer=AsyncMock()))
        await ui.cash_callback(cb(self.manager,'cash:issue'))
        await ui.cash_callback(cb(self.manager,f'cash:employee:{employee}'))
        await ui.cash_callback(cb(self.manager,f'cash:project:{employee}:{self.project.id}'))
        message=SimpleNamespace(from_user=self.manager,text='500 Закупка краски',answer=AsyncMock())
        self.assertTrue(await ui.finance_input(message))
        await ui.amount_input(message)
        aid=(await ui.service.list(manager))[0]['id']
        await ui.cash_callback(cb(self.employee,f'cash:receive:{aid}'))
        self.assertIsNotNone((await ui.service.card(employee,aid))['received_at'])
        await ui.list_advances(SimpleNamespace(from_user=self.employee,answer=AsyncMock()))
        message.text='Проекты'
        self.assertFalse(await ui.finance_input(message))

    async def test_advance_cannot_issue_against_foreign_company_project(self):
        from app.services.advances import AdvanceService,AdvanceError
        manager,employee=await self.finance_ids()
        other=await self.service.create_company(self.owner,'Other company')
        async with self.pool.acquire() as c:
            pid=await c.fetchval('INSERT INTO projects(company_id,name,created_by_user_id) SELECT $1,$2,id FROM users WHERE telegram_id=$3 RETURNING id',other.id,'Foreign',self.owner.id)
        with self.assertRaises(AdvanceError):
            await AdvanceService().issue(manager,employee,pid,100,'x','foreign')

    async def test_issue_explains_empty_employee_list(self):
        from app.handlers import advances as ui
        async with self.pool.acquire() as c:
            await c.execute("UPDATE company_members SET status='removed' WHERE role='employee'")
        callback=SimpleNamespace(id='empty',from_user=self.manager,data='cash:issue',answer=AsyncMock(),message=SimpleNamespace(answer=AsyncMock()))
        await ui.cash_callback(callback)
        self.assertIn('нет активных сотрудников',callback.message.answer.call_args.args[0])
        self.assertNotIn('reply_markup',callback.message.answer.call_args.kwargs)

    async def test_money_on_hand_posts_expenses_but_waits_for_cash_confirmation(self):
        from app.services.advances import AdvanceService
        s=AdvanceService();manager,employee=await self.finance_ids()
        aid=await s.issue(manager,employee,self.project.id,200,'x','received')
        await s.issue(manager,employee,self.project.id,300,'x','pending')
        row=(await s.balances(employee))[0]
        self.assertEqual(row['balance'],0);self.assertEqual(row['pending'],500)
        await s.receive(employee,aid)
        doc=await DocumentService().save_document(self.employee,self.project,'preview',self.document,upload_id='balance-doc')
        eid=await s.attach(employee,aid,doc)
        rid=await s.request_return(employee,aid,25,'return-balance')
        self.assertEqual((await s.balances(employee))[0]['balance'],100)
        await s.approve_expense(manager,aid,eid);await s.approve_return(manager,aid,rid)
        row=(await s.balances(employee))[0]
        self.assertEqual(row['balance'],75);self.assertEqual(row['pending'],300)
        self.assertEqual((await s.balances(manager))[0]['balance'],75)

    async def test_money_on_hand_is_personal_and_sums_across_projects(self):
        from app.services.advances import AdvanceService
        s=AdvanceService();manager,employee=await self.finance_ids()
        project=await ProjectService().create_project(self.manager.id,'Second project')
        for key,pid,amount in [('one',self.project.id,50),('two',project.id,70)]:
            aid=await s.issue(manager,employee,pid,amount,'x',key);await s.receive(employee,aid)
        self.assertEqual((await s.balances(employee))[0]['balance'],120)
        outsider=await self.service.ensure_platform_user(self.user(1998))
        await self.service.assign_user_to_company_by_owner(self.owner.id,outsider,self.company.id,'employee')
        self.assertEqual(await s.balances(outsider),[])

    async def test_balance_menu_does_not_consume_amount_input(self):
        from app.handlers import advances as ui
        await set_pending_action(self.employee.id,'advance_return',{'advance':1,'key':'x'})
        message=SimpleNamespace(from_user=self.employee,text=ui.BALANCE_MENU,answer=AsyncMock())
        self.assertFalse(await ui.finance_input(message))
        await ui.balances_entry(message)
        self.assertIn('0.00 RUB',message.answer.call_args.args[0])

    async def test_money_report_opening_and_approval_dates_and_pending(self):
        from datetime import datetime,timedelta,timezone
        from app.services.advances import AdvanceService
        s=AdvanceService();manager,employee=await self.finance_ids()
        aid=await s.issue(manager,employee,self.project.id,200,'x','report')
        await s.receive(employee,aid)
        await s.issue(manager,employee,self.project.id,300,'x','pending-report')
        doc=await DocumentService().save_document(self.employee,self.project,'preview',self.document,upload_id='report-doc')
        eid=await s.attach(employee,aid,doc)
        await self.legacy_pending_expense(eid)
        rid=await s.request_return(employee,aid,25,'report-return')
        now=datetime.now(timezone.utc);start=now-timedelta(days=1)
        async with self.pool.acquire() as c:
            await c.execute('UPDATE advances SET created_at=$2,received_at=$2 WHERE id=$1',aid,now-timedelta(days=2))
        before=await s.report(manager,start)
        self.assertEqual(before['totals']['opening'],200)
        self.assertEqual(before['totals']['received'],0)
        self.assertEqual(before['totals']['closing'],200)
        self.assertEqual(before['totals']['pending'],300)
        self.assertEqual(before['totals']['pending_expenses'],100)
        self.assertEqual(before['totals']['pending_returns'],25)
        await s.approve_expense(manager,aid,eid);await s.approve_return(manager,aid,rid)
        after=await s.report(manager,start)
        self.assertEqual(after['totals']['spent'],100)
        self.assertEqual(after['totals']['returned'],25)
        self.assertEqual(after['totals']['closing'],75)
        historical=await s.report(manager,start,now)
        self.assertEqual(historical['totals']['closing'],200)
        self.assertEqual(after['totals']['pending_expenses'],0)
        self.assertEqual(after['totals']['pending_returns'],0)

    async def test_money_report_scope_and_excel(self):
        from datetime import datetime,timezone
        from app.services.advances import AdvanceService,AdvanceError
        from app.services.report_exports import build_manager_report_workbook
        from openpyxl import load_workbook
        from io import BytesIO
        s=AdvanceService();manager,employee=await self.finance_ids()
        aid=await s.issue(manager,employee,self.project.id,150,'x','report')
        await s.receive(employee,aid)
        start=datetime(2000,1,1,tzinfo=timezone.utc)
        report=await s.report(manager,start,project_id=self.project.id)
        self.assertEqual(report['totals']['closing'],150)
        self.assertEqual((await s.report(manager,start,project_id=-1))['rows'],[])
        with self.assertRaises(AdvanceError): await s.report(employee,start,employee_id=manager)
        summary=SimpleNamespace(start_at=start,documents=0,total_amount=0,exact_duplicates=0,probable_duplicates=0)
        _,data=build_manager_report_workbook('all_time',summary,[],[],[],[],[],money_report=report)
        wb=load_workbook(BytesIO(data))
        self.assertIn('Money',wb.sheetnames)
        self.assertEqual(wb['Money'].cell(2,9).value,150)
        self.assertEqual(wb['Money'].max_row,3)

    async def test_money_report_callback_outputs_financial_summary(self):
        from app.handlers.start import report_period_callback
        from app.ui.reports import MANAGER_REPORTS_PERIOD_PREFIX
        from app.services.advances import AdvanceService
        manager,employee=await self.finance_ids()
        aid=await AdvanceService().issue(manager,employee,self.project.id,50,'x','ui-report')
        await AdvanceService().receive(employee,aid)
        callback=SimpleNamespace(from_user=self.manager,data=MANAGER_REPORTS_PERIOD_PREFIX+'money:all_time',answer=AsyncMock(),message=SimpleNamespace(answer=AsyncMock()))
        await report_period_callback(callback)
        text=callback.message.answer.call_args.args[0]
        self.assertIn('Отчёт по деньгам',text)
        self.assertIn('Получено: 50.00',text)
        self.assertIn('Остаток',text.replace('остаток','Остаток'))

    async def test_business_card_payment_approval_does_not_change_cash_balance(self):
        from app.services.business_cards import BusinessCardService
        from app.services.advances import AdvanceService,AdvanceError
        from datetime import datetime,timezone
        s=BusinessCardService();manager,employee=await self.finance_ids()
        card=await s.register(manager,'Основная','0012')
        self.assertEqual(await s.register(manager,'Основная','0012'),card)
        doc=await DocumentService().save_document(self.employee,self.project,'preview',self.document,upload_id='card-doc')
        eid=await s.submit(employee,card,doc)
        self.assertEqual(await s.submit(employee,card,doc),eid)
        with self.assertRaises(AdvanceError): await s.approve(employee,eid)
        await s.approve(manager,eid);await s.approve(manager,eid)
        self.assertEqual(await AdvanceService().balances(employee),[])
        rows=await s.report(manager,datetime(2000,1,1,tzinfo=timezone.utc))
        self.assertEqual(len(rows),1);self.assertTrue(rows[0]['accepted']);self.assertEqual(rows[0]['last4'],'0012')

    async def test_card_and_cash_cannot_claim_same_document_concurrently(self):
        from app.services.business_cards import BusinessCardService
        from app.services.advances import AdvanceService,AdvanceError
        cards=BusinessCardService();cash=AdvanceService();manager,employee=await self.finance_ids()
        card=await cards.register(manager,'Company','1234')
        aid=await cash.issue(manager,employee,self.project.id,200,'x','card-race');await cash.receive(employee,aid)
        doc=await DocumentService().save_document(self.employee,self.project,'preview',self.document,upload_id='card-race')
        result=await asyncio.gather(cards.submit(employee,card,doc),cash.attach(employee,aid,doc),return_exceptions=True)
        self.assertEqual(sum(isinstance(r,AdvanceError) for r in result),1)
        async with self.pool.acquire() as c:
            self.assertEqual(await c.fetchval('SELECT (SELECT count(*) FROM card_expenses)+(SELECT count(*) FROM advance_expenses)'),1)

    async def test_card_permissions_and_invalid_numbers(self):
        from app.services.business_cards import BusinessCardService
        from app.services.advances import AdvanceError
        s=BusinessCardService();manager,employee=await self.finance_ids()
        with self.assertRaises(AdvanceError): await s.register(employee,'Forbidden','1234')
        for last4 in ('123','12345','1234567890123456','abcd'):
            with self.assertRaises(AdvanceError): await s.register(manager,'Card',last4)
        with self.assertRaises(AdvanceError): await s.register(manager,'Card 1234567890123456','1234')
        card=await s.register(manager,'Valid','1234')
        doc=await DocumentService().save_document(self.employee,self.project,'preview',self.document,upload_id='owner-card')
        with self.assertRaises(AdvanceError): await s.submit(manager,card,doc)
        outsider=await self.service.ensure_platform_user(self.user(1777))
        with self.assertRaises(AdvanceError): await s.cards(outsider)

    async def test_card_report_excel_and_personal_visibility(self):
        from app.services.business_cards import BusinessCardService
        from app.handlers.money_reports import get_money_report
        from app.services.money_reports import format_money_report
        from app.services.report_exports import build_manager_report_workbook
        from datetime import datetime,timezone,timedelta
        from openpyxl import load_workbook
        from io import BytesIO
        s=BusinessCardService();manager,employee=await self.finance_ids()
        card=await s.register(manager,'=1+1','0012')
        doc=await DocumentService().save_document(self.employee,self.project,'preview',self.document,upload_id='card-xlsx')
        eid=await s.submit(employee,card,doc)
        before=await get_money_report(self.manager.id,'all_time')
        self.assertEqual(before['card_pending'],0);self.assertEqual(before['card_spent'],100)
        await s.approve(manager,eid)
        report=await get_money_report(self.manager.id,'all_time')
        self.assertEqual(report['card_spent'],100);self.assertEqual(report['totals']['spent'],0)
        self.assertIn('Бизнес-карты',format_money_report(report))
        summary=SimpleNamespace(start_at=datetime(2000,1,1,tzinfo=timezone.utc),documents=1,total_amount=100,exact_duplicates=0,probable_duplicates=0)
        _,payload=build_manager_report_workbook('all_time',summary,[],[],[],[],[],money_report=report)
        workbook=load_workbook(BytesIO(payload));sheet=workbook['BusinessCards']
        self.assertEqual(sheet.cell(2,2).value,'=1+1');self.assertEqual(sheet.cell(2,2).data_type,'s')
        self.assertEqual(sheet.cell(2,3).value,'0012')
        self.assertEqual(await s.report(manager,datetime.now(timezone.utc)+timedelta(days=1),datetime.now(timezone.utc)+timedelta(days=2)),[])

    async def test_business_cards_telegram_registration_and_payment(self):
        from app.handlers import business_cards as ui
        from app.services.business_cards import BusinessCardService
        def cb(user,data):
            return SimpleNamespace(from_user=user,data=data,answer=AsyncMock(),message=SimpleNamespace(answer=AsyncMock()))
        await ui.card_callback(cb(self.manager,'bcard:new'))
        message=SimpleNamespace(from_user=self.manager,text='Основная 9876',answer=AsyncMock())
        self.assertTrue(await ui.card_input(message))
        await ui.create_card(message)
        manager,employee=await self.finance_ids()
        card=(await BusinessCardService().cards(employee))[0]['id']
        doc=await DocumentService().save_document(self.employee,self.project,'preview',self.document,upload_id='card-ui')
        await ui.card_callback(cb(self.employee,'bcard:documents'))
        await ui.card_callback(cb(self.employee,f'bcard:document:{doc}'))
        await ui.card_callback(cb(self.employee,f'bcard:submit:{doc}:{card}'))
        await ui.cards_entry(SimpleNamespace(from_user=self.manager,answer=AsyncMock()))
        async with self.pool.acquire() as c:
            self.assertEqual(await c.fetchval('SELECT count(*) FROM card_expenses'),1)

    async def test_account_link_requires_original_confirmation_and_preserves_membership(self):
        from app.services.account_links import AccountLinkService
        s=AccountLinkService()
        request=await s.issue('telegram',str(self.employee.id),'max')
        rid=await s.claim('max','max-123',request['code'])
        self.assertIsNone(await s.resolve('max','max-123'))
        pending=await s.pending('telegram',str(self.employee.id))
        self.assertEqual(pending[0]['id'],rid)
        uid=await s.confirm('telegram',str(self.employee.id),rid)
        self.assertEqual(await s.resolve('max','max-123'),uid)
        self.assertEqual(await s.confirm('telegram',str(self.employee.id),rid),uid)
        self.assertEqual((await self.service.get_user_context(self.employee.id)).member_role,'employee')
        async with self.pool.acquire() as c:
            row=await c.fetchrow('SELECT code_hash FROM account_link_requests WHERE id=$1',rid)
            self.assertNotEqual(row['code_hash'],request['code'])

    async def test_account_link_expiration_replay_and_other_user_confirmation_denied(self):
        from app.services.account_links import AccountLinkService,AccountLinkError
        s=AccountLinkService();request=await s.issue('telegram',str(self.employee.id),'max')
        rid=await s.claim('max','max-123',request['code'])
        with self.assertRaises(AccountLinkError): await s.claim('max','max-456',request['code'])
        with self.assertRaises(AccountLinkError): await s.confirm('telegram',str(self.manager.id),rid)
        async with self.pool.acquire() as c:
            await c.execute("UPDATE account_link_requests SET expires_at=NOW()-interval '1 second' WHERE id=$1",rid)
        with self.assertRaises(AccountLinkError): await s.confirm('telegram',str(self.employee.id),rid)
        self.assertIsNone(await s.resolve('max','max-123'))

    async def test_account_link_invalid_attempts_are_counted_and_throttled(self):
        from app.services.account_links import AccountLinkService,AccountLinkError
        s=AccountLinkService()
        for _ in range(5):
            with self.assertRaises(AccountLinkError): await s.claim('max','guessing-user','wrong')
        request=await s.issue('telegram',str(self.employee.id),'max')
        with self.assertRaisesRegex(AccountLinkError,'Лимит'):
            await s.claim('max','guessing-user',request['code'])
        async with self.pool.acquire() as c:
            self.assertEqual(await c.fetchval("SELECT count(*) FROM account_link_attempts WHERE external_id='guessing-user'"),5)

    async def test_account_link_does_not_merge_existing_accounts_and_new_code_revokes_old(self):
        from app.services.account_links import AccountLinkService,AccountLinkError
        s=AccountLinkService()
        first=await s.issue('telegram',str(self.employee.id),'max')
        second=await s.issue('telegram',str(self.employee.id),'max')
        with self.assertRaises(AccountLinkError): await s.claim('max','one',first['code'])
        rid=await s.claim('max','one',second['code'])
        await s.confirm('telegram',str(self.employee.id),rid)
        manager_request=await s.issue('telegram',str(self.manager.id),'max')
        with self.assertRaises(AccountLinkError): await s.claim('max','one',manager_request['code'])
        with self.assertRaises(AccountLinkError): await s.issue('telegram',str(self.employee.id),'max')

    async def test_concurrent_claim_of_code_has_only_one_winner(self):
        from app.services.account_links import AccountLinkService,AccountLinkError
        s=AccountLinkService();request=await s.issue('telegram',str(self.employee.id),'max')
        results=await asyncio.gather(s.claim('max','one',request['code']),s.claim('max','two',request['code']),return_exceptions=True)
        self.assertEqual(sum(isinstance(x,AccountLinkError) for x in results),1)
        self.assertIsNone(await s.resolve('max','one'));self.assertIsNone(await s.resolve('max','two'))

    async def test_link_notification_is_sent_once_to_original_account(self):
        from app.services.account_links import AccountLinkService
        from app.handlers.link_notifications import deliver_once
        s=AccountLinkService();request=await s.issue('telegram',str(self.employee.id),'max')
        await s.claim('max','candidate',request['code'])
        bot=SimpleNamespace(send_message=AsyncMock())
        await deliver_once(bot);await deliver_once(bot)
        self.assertEqual(bot.send_message.await_count,1)
        self.assertEqual(bot.send_message.call_args.kwargs['chat_id'],self.employee.id)
        self.assertIsNone(await s.resolve('max','candidate'))

    async def test_cancel_link_and_inactive_user_cannot_complete(self):
        from app.services.account_links import AccountLinkService,AccountLinkError
        s=AccountLinkService();request=await s.issue('telegram',str(self.employee.id),'max')
        await s.cancel('telegram',str(self.employee.id),request['id'])
        with self.assertRaises(AccountLinkError): await s.claim('max','candidate',request['code'])
        request=await s.issue('telegram',str(self.employee.id),'max')
        await s.claim('max','candidate',request['code'])
        async with self.pool.acquire() as c:
            await c.execute('UPDATE users SET is_active=FALSE WHERE telegram_id=$1',self.employee.id)
        with self.assertRaises(AccountLinkError): await s.confirm('telegram',str(self.employee.id),request['id'])

    async def test_linking_telegram_to_existing_identity_does_not_create_duplicate_user(self):
        from app.services.account_links import AccountLinkService
        from app.handlers.profile import claim_link
        from aiogram.filters import CommandObject
        s=AccountLinkService()
        async with self.pool.acquire() as c:
            uid=await c.fetchval("INSERT INTO users(first_name) VALUES('Existing MAX user') RETURNING id")
            await c.execute("INSERT INTO external_accounts(user_id,platform,external_id) VALUES($1,'max','existing-max')",uid)
            count=await c.fetchval('SELECT count(*) FROM users')
        request=await s.issue('max','existing-max','telegram')
        new_user=self.user(1987654)
        message=SimpleNamespace(from_user=new_user,chat=SimpleNamespace(type='private'),answer=AsyncMock())
        await claim_link(message,CommandObject(prefix='/',command='link',args=request['code']))
        self.assertIsNone(await s.resolve('telegram',str(new_user.id)))
        await s.confirm('max','existing-max',request['id'])
        self.assertEqual(await self.service.ensure_platform_user(new_user),uid)
        async with self.pool.acquire() as c:
            self.assertEqual(await c.fetchval('SELECT count(*) FROM users'),count)

    async def test_profile_rejects_group_chat_without_disclosing_code(self):
        from app.handlers.profile import profile_entry
        message=SimpleNamespace(from_user=self.employee,chat=SimpleNamespace(type='group'),answer=AsyncMock())
        await profile_entry(message)
        self.assertIn('личном чате',message.answer.call_args.args[0])
        async with self.pool.acquire() as c:
            self.assertEqual(await c.fetchval('SELECT count(*) FROM account_link_requests'),0)
