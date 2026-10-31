import os
import unittest
from datetime import datetime, timezone, date
from io import BytesIO
from unittest.mock import AsyncMock, patch
from zipfile import ZipFile

import httpx
from openpyxl import load_workbook
from db_fixture import DatabaseFixture
from test_document_core import photo_bytes
from test_max_adapter import message
from app.application.accountant import AccountantCommands
from app.max_bot.handler import MaxHandler
from app.services.document_flow import DocumentFlow
from app.services.accountant_exports import AccountantExportService, accountant_dates
from app.services.advances import AdvanceError
from app.services.workspace import WorkspaceService
from app.web.app import create_app, manager as web_manager

class AccountantPeriodTests(unittest.TestCase):
    def test_presets_custom_and_invalid_ranges(self):
        now = datetime(2026, 1, 15, tzinfo=timezone.utc)
        self.assertEqual(accountant_dates('previous_month',now=now),(date(2025,12,1),date(2025,12,31)))
        self.assertEqual(accountant_dates('2026-09-01 2026-10-31'),(date(2026,9,1),date(2026,10,31)))
        self.assertEqual(accountant_dates('all_time',now=now)[0],date.min)
        for value in ('bad','2026-10-31 2026-09-01','2026-02-30 2026-03-01','2026-01-01 9999-12-31'):
            with self.assertRaises(AdvanceError): accountant_dates(value)

@unittest.skipUnless(os.environ.get('RUN_DATABASE_TESTS') == '1','Disposable PostgreSQL required')
class AccountantExportTests(DatabaseFixture):
    async def asyncSetUp(self):
        await super().asyncSetUp()
        self.manager_id, self.employee_id = await self.finance_ids()
        self.flow = DocumentFlow()
        self.export = AccountantExportService()
        self.recognition = patch('app.services.document_flow.recognize_document',AsyncMock(return_value=('text', self.document)))
        self.recognition.start()
        self.addCleanup(self.recognition.stop)

    async def receive(self, source='photo', data=None):
        return await self.flow.receive(self.employee_id,'max',os.urandom(16).hex(),source,AsyncMock(return_value=data or photo_bytes()))

    async def save(self, source='photo', data=None):
        view = await self.receive(source,data)
        return await self.flow.action(self.employee_id,'project',view.upload_id,str(self.project.id))

    async def test_exact_original_and_positions_and_literal_names(self):
        data=photo_bytes()
        await self.pool.execute("UPDATE projects SET name='=HYPERLINK(\"https://evil.invalid\")' WHERE id=$1",self.project.id)
        saved=await self.save(data=data)
        self.assertEqual(await self.pool.fetchval('SELECT content FROM document_originals'), data)
        name,payload,missing=await self.export.build(self.manager_id,'all_time')
        self.assertTrue(name.endswith('.zip'))
        self.assertEqual(missing,0)
        with ZipFile(BytesIO(payload)) as archive:
            self.assertEqual(archive.read(f'files/{saved.document_id}.jpg'),data)
            book=load_workbook(BytesIO(archive.read('Реестр.xlsx')))
            self.assertEqual(book['Документы']['B2'].data_type,'s')
            self.assertEqual(book['Позиции'].max_row,2)
        self.assertEqual(await self.pool.fetchval('SELECT count(*) FROM pending_document_flows'),0)
        await self.flow.action(self.employee_id,'project',saved.upload_id,str(self.project.id))
        self.assertEqual(await self.pool.fetchval('SELECT count(*) FROM document_originals'),1)

    async def test_pdf_kept_as_pdf_and_inclusive_document_date(self):
        data=b'%PDF-test-original'
        saved=await self.save('pdf',data)
        await self.pool.execute("UPDATE documents SET document_date='2026-09-30 23:59:59+00',created_at='2026-10-20' WHERE id=$1", saved.document_id)
        _, payload, _ = await self.export.build(self.manager_id,'2026-09-01 2026-09-30')
        with ZipFile(BytesIO(payload)) as archive:
            self.assertEqual(archive.read(f'files/{saved.document_id}.pdf'),data)
        with self.assertRaises(AdvanceError): await self.export.build(self.manager_id,'2026-10-01 2026-10-31')

    async def test_old_missing_scan_is_listed_with_fallback_date_archived_project(self):
        from app.adapters.telegram.documents import DocumentService
        await DocumentService().save_document(self.employee,self.project,'text',self.document)
        await self.pool.execute("UPDATE documents SET document_date=NULL,created_at='2026-09-30 23:59:59+00'")
        await self.pool.execute("UPDATE projects SET status='archived' WHERE id=$1",self.project.id)
        _,payload,missing=await self.export.build(self.manager_id,'2026-09-01 2026-09-30')
        self.assertEqual(missing,1)
        with ZipFile(BytesIO(payload)) as archive:
            self.assertEqual(archive.namelist(),['Реестр.xlsx'])
            book=load_workbook(BytesIO(archive.read('Реестр.xlsx')))
            self.assertEqual(book['Документы']['N2'].value,'Скан отсутствует')
            self.assertEqual(book['Документы']['O2'].value,'Дата загрузки UTC')

    async def test_employee_and_inactive_manager_and_foreign_project_denied(self):
        await self.save()
        with self.assertRaises(AdvanceError): await self.export.build(self.employee_id,'all_time')
        with self.assertRaises(AdvanceError): await self.export.build(self.manager_id,'all_time',project_id=999999)
        await self.pool.execute('UPDATE users SET is_active=false WHERE id=$1',self.manager_id)
        with self.assertRaises(AdvanceError): await self.export.build(self.manager_id,'all_time')

    async def test_original_insert_failure_rolls_back_document_and_retains_draft(self):
        view=await self.receive()
        await self.pool.execute("ALTER TABLE document_originals ADD CONSTRAINT fail_test CHECK(false)")
        import asyncpg
        with self.assertRaises(asyncpg.CheckViolationError):
            await self.flow.action(self.employee_id,'project',view.upload_id,str(self.project.id))
        self.assertEqual(await self.document_count(),0)
        self.assertIsNotNone(await self.pool.fetchval('SELECT original_data FROM pending_document_flows'))
        await self.pool.execute('ALTER TABLE document_originals DROP CONSTRAINT fail_test')
        await self.flow.action(self.employee_id,'project',view.upload_id,str(self.project.id))
        self.assertEqual(await self.document_count(),1)

    async def test_cancel_purges_pending_scan_and_size_limits_do_not_truncate(self):
        view=await self.receive()
        await self.flow.action(self.employee_id,'cancel',view.upload_id)
        self.assertEqual(await self.pool.fetchval('SELECT count(*) FROM pending_document_flows'),0)
        await self.save()
        with patch('app.services.accountant_exports.MAX_SOURCE_BYTES',1):
            with self.assertRaisesRegex(AdvanceError,'короткий'): await self.export.build(self.manager_id,'all_time')
        with patch('app.services.accountant_exports.MAX_ARCHIVE_BYTES',1):
            with self.assertRaisesRegex(AdvanceError,'45 МБ'): await self.export.build(self.manager_id,'all_time')

    async def test_access_rechecked_after_packaging(self):
        await self.save()
        with patch.object(self.export,'authorize',AsyncMock(return_value=-1)):
            with self.assertRaisesRegex(AdvanceError,'Компания изменилась'): await self.export.build(self.manager_id,'all_time')

    async def test_max_menu_custom_dates_and_panel_retains_selected_range(self):
        await self.pool.execute("INSERT INTO external_accounts(user_id,platform,external_id) VALUES($1,'max','9001')",self.manager_id)
        handler=MaxHandler()
        choice=(await handler.handle(message(9001,'/accountant')))[1]
        self.assertIn('Прошлый месяц',[str(r[0]) for r in choice.buttons])
        self.assertIn('Укажите период',await AccountantCommands().selection(self.manager_id,'custom'))
        with patch.dict(os.environ,{'WEB_ORIGIN':'https://testserver'}):
            result=(await handler.handle(message(9001,'/accountant 2026-09-01 2026-10-31')))[1]
            self.assertIn('start=2026-09-01&end=2026-10-31',result.buttons[0][0].url)

    async def test_http_export_authenticated_only_and_exact_attachment(self):
        saved=await self.save()
        app=create_app()
        async with httpx.AsyncClient(transport=httpx.ASGITransport(app=app),base_url='https://testserver') as client:
            url='/api/reports/accountant?start=2026-01-01&end=2026-12-31'
            self.assertEqual((await client.get(url)).status_code,401)
            app.dependency_overrides[web_manager]=lambda: {'id':self.manager_id}
            response=await client.get(url)
            self.assertEqual(response.status_code,200,response.text[:300] if response.status_code!=200 else '')
            self.assertEqual(response.headers['content-type'],'application/zip')
            self.assertEqual(response.headers['x-missing-scans'],'0')
            with ZipFile(BytesIO(response.content)) as archive:
                self.assertIn(f'files/{saved.document_id}.jpg',archive.namelist())
            self.assertEqual((await client.get('/api/reports/accountant?start=bad&end=2026-12-31')).status_code,422)

    async def test_other_company_cannot_export_or_filter_our_original(self):
        await self.save()
        other = await self.service.create_company(self.owner, 'Other company')
        invitation = await self.service.create_initial_manager_invite(self.owner, other.id)
        person = self.user(1099)
        await self.service.join_company(person, invitation)
        other_actor = await self.pool.fetchval('SELECT id FROM users WHERE telegram_id=$1', person.id)
        with self.assertRaisesRegex(AdvanceError, 'документов нет'):
            await self.export.build(other_actor, 'all_time')
        with self.assertRaisesRegex(AdvanceError, 'недоступен'):
            await self.export.build(other_actor, 'all_time',project_id=self.project.id)

    async def test_telegram_delivery_and_period_buttons_call_shared_export(self):
        from types import SimpleNamespace
        from app.handlers.accountant import show
        await self.save()
        outgoing = SimpleNamespace(answer=AsyncMock(),answer_document=AsyncMock())
        await show(outgoing,self.manager,'/accountant')
        labels=[button.text for row in outgoing.answer.call_args.kwargs['reply_markup'].inline_keyboard for button in row]
        self.assertIn('Указать даты',labels)
        await show(outgoing,self.manager,'/accountant all_time')
        file=outgoing.answer_document.call_args.args[0]
        with ZipFile(BytesIO(file.data)) as archive:
            self.assertEqual(len(archive.namelist()),2)
        self.assertIn('Без скана: 0',outgoing.answer_document.call_args.kwargs['caption'])

    async def test_zero_values_remain_distinct_from_unknown_in_accountant_register(self):
        await self.save()
        await self.pool.execute('UPDATE document_items SET quantity=NULL,price=0,line_total=0')
        _, payload, _ = await self.export.build(self.manager_id, 'all_time')
        with ZipFile(BytesIO(payload)) as archive:
            book = load_workbook(BytesIO(archive.read('Реестр.xlsx')))
        from decimal import Decimal
        self.assertIsNone(book['Позиции']['D2'].value)
        self.assertEqual(Decimal(book['Позиции']['E2'].value), 0)
        self.assertEqual(Decimal(book['Позиции']['F2'].value), 0)

    async def test_navigation_never_fetches_original_and_save_transfers_it_inside_database(self):
        data = photo_bytes()
        view = await self.receive(data=data)
        async with self.flow.transaction(self.employee_id) as (connection, member):
            row = await self.flow._row(connection, self.employee_id, member, view.upload_id)
            self.assertNotIn('original_data', row.keys())
            self.assertNotIn('content', row.keys())
        resumed = await self.flow.action(self.employee_id, 'projects', view.upload_id, '1')
        self.assertEqual(resumed.upload_id, view.upload_id)
        saved = await self.flow.action(self.employee_id, 'project', view.upload_id, str(self.project.id))
        self.assertEqual(await self.pool.fetchval('SELECT content FROM document_originals WHERE document_id=$1', saved.document_id), data)

    async def test_originals_are_loaded_in_bounded_batches_and_temporary_files_are_removed(self):
        import asyncpg
        import tempfile
        from pathlib import Path
        from app.services import accountant_exports as exports
        saved = await self.save()
        original = await self.pool.fetchrow('SELECT content,extension FROM document_originals WHERE document_id=$1', saved.document_id)
        for _ in range(4):
            clone = await self.flow.documents.save_for_user(self.employee_id, self.project.id, 'preview', self.document)
            await self.pool.execute('INSERT INTO document_originals(document_id,content,extension) VALUES($1,$2,$3)', clone, original['content'], original['extension'])
        batches = []
        fetch = asyncpg.Connection.fetch
        async def observed(connection, query, *args, **kwargs):
            if 'SELECT document_id,content FROM document_originals' in query:
                batches.append(tuple(args[0]))
            return await fetch(connection, query, *args, **kwargs)
        with tempfile.TemporaryDirectory() as directory, patch.object(exports, 'TMP_DIR', Path(directory)), patch.object(exports, 'ORIGINAL_BATCH_COUNT', 2), patch.object(asyncpg.Connection, 'fetch', new=observed):
            _, payload, missing = await self.export.build(self.manager_id, 'all_time')
            self.assertEqual(list(Path(directory).iterdir()), [])
        self.assertEqual([len(batch) for batch in batches], [2,2,1])
        self.assertEqual(missing, 0)
        with ZipFile(BytesIO(payload)) as archive:
            self.assertEqual(len(archive.namelist()), 6)
            self.assertEqual(archive.read(f'files/{saved.document_id}.jpg'), original['content'])

    async def test_expired_scan_cleanup_is_bounded_and_keeps_active_or_locked_drafts(self):
        from app.services.document_housekeeping import cleanup_expired_document_drafts
        actors = [self.employee_id, self.manager_id]
        for actor in actors:
            await self.pool.execute("INSERT INTO pending_document_flows(actor_id,company_id,upload_id,expires_at,original_data,original_ext) VALUES($1,$2,$3,clock_timestamp()-interval '1 minute',$4,'jpg')", actor, self.company.id, f'expired-{actor}', photo_bytes())
        async with self.pool.acquire() as connection, connection.transaction():
            await connection.execute('SELECT actor_id FROM pending_document_flows WHERE actor_id=$1 FOR UPDATE', self.employee_id)
            self.assertEqual(await cleanup_expired_document_drafts(limit=1), 1)
            self.assertIsNotNone(await connection.fetchval('SELECT original_data FROM pending_document_flows WHERE actor_id=$1', self.employee_id))
        await self.pool.execute("UPDATE pending_document_flows SET expires_at=clock_timestamp()+interval '1 hour' WHERE actor_id=$1", self.employee_id)
        self.assertEqual(await cleanup_expired_document_drafts(), 0)
        await self.pool.execute("UPDATE pending_document_flows SET expires_at=clock_timestamp()-interval '1 minute' WHERE actor_id=$1", self.employee_id)
        self.assertEqual(await cleanup_expired_document_drafts(), 1)
        self.assertEqual(await self.pool.fetchval('SELECT count(*) FROM pending_document_flows'), 0)


    async def test_waiting_download_rechecks_revoked_access_before_reading_bytes(self):
        import asyncio
        from types import SimpleNamespace
        from app.services.processing_limits import document_processing_slot
        from app.services.membership import MembershipError
        download = AsyncMock(return_value=photo_bytes())
        with patch('app.services.processing_limits.get_recognition_settings', return_value=SimpleNamespace(document_processing_concurrency=1)):
            async with document_processing_slot():
                task = asyncio.create_task(self.flow.receive(self.employee_id, 'max', 'waiting-access', 'photo', download))
                try:
                    async with asyncio.timeout(2):
                        while not await self.pool.fetchval('SELECT count(*) FROM pending_document_flows WHERE actor_id=$1', self.employee_id):
                            await asyncio.sleep(.005)
                    download.assert_not_awaited()
                    await self.pool.execute('UPDATE users SET is_active=false WHERE id=$1', self.employee_id)
                except BaseException:
                    task.cancel()
                    await asyncio.gather(task, return_exceptions=True)
                    raise
            with self.assertRaises(MembershipError):
                await asyncio.wait_for(task, 2)
        download.assert_not_awaited()
        self.assertEqual(await self.pool.fetchval('SELECT count(*) FROM pending_document_flows WHERE actor_id=$1', self.employee_id), 0)

    async def test_cancelled_waiting_upload_clears_draft_without_downloading(self):
        import asyncio
        from types import SimpleNamespace
        from app.services.processing_limits import document_processing_slot
        download = AsyncMock(return_value=photo_bytes())
        with patch('app.services.processing_limits.get_recognition_settings', return_value=SimpleNamespace(document_processing_concurrency=1)):
            async with document_processing_slot():
                task = asyncio.create_task(self.flow.receive(self.employee_id, 'max', 'waiting-cancel', 'photo', download))
                try:
                    async with asyncio.timeout(2):
                        while not await self.pool.fetchval('SELECT count(*) FROM pending_document_flows WHERE actor_id=$1', self.employee_id):
                            await asyncio.sleep(.005)
                    download.assert_not_awaited()
                finally:
                    task.cancel()
                    await asyncio.gather(task, return_exceptions=True)
        self.assertTrue(task.cancelled())
        download.assert_not_awaited()
        self.assertEqual(await self.pool.fetchval('SELECT count(*) FROM pending_document_flows WHERE actor_id=$1', self.employee_id), 0)


class AccountantWorkerTests(unittest.IsolatedAsyncioTestCase):
    async def test_batch_size_limit_and_oversized_single_file(self):
        from app.services.accountant_exports import _original_batches
        rows = [{'id':i, 'original_size':size} for i,size in enumerate((4,4,20,0,1),1)]
        with patch('app.services.accountant_exports.ORIGINAL_BATCH_BYTES', 8):
            self.assertEqual(list(_original_batches(rows)), [[1,2],[3],[5]])

    async def test_housekeeping_task_stops_before_database_shutdown(self):
        import asyncio
        from app.services import document_housekeeping as housekeeping
        with patch.object(housekeeping, 'cleanup_expired_document_drafts', AsyncMock(return_value=0)) as cleanup:
            task = housekeeping.start_document_housekeeping()
            await asyncio.sleep(0)
            await housekeeping.stop_document_housekeeping(task)
        cleanup.assert_awaited_once()
        self.assertTrue(task.done())

    async def test_cancelled_file_work_finishes_before_unwinding_owned_files(self):
        import asyncio
        import threading
        from app.services.accountant_exports import _file_work
        started, finish, finished = threading.Event(), threading.Event(), threading.Event()
        def write():
            started.set()
            finish.wait(2)
            finished.set()
        task = asyncio.create_task(_file_work(write))
        while not started.is_set():
            await asyncio.sleep(.001)
        task.cancel()
        await asyncio.sleep(0)
        self.assertFalse(task.done())
        task.cancel()  # Shutdown can cancel both a worker and its owning poller.
        await asyncio.sleep(0)
        self.assertFalse(task.done())
        finish.set()
        with self.assertRaises(asyncio.CancelledError):
            await task
        self.assertTrue(finished.is_set())
