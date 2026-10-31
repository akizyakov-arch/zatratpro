import os
import unittest
from io import BytesIO
from zipfile import ZipFile
from decimal import Decimal
from datetime import datetime,timezone
from openpyxl import load_workbook
from db_fixture import DatabaseFixture
from app.schemas.document import DocumentSchema,DocumentItem
from app.schemas.vat import document_vat_total
from app.services.json_formatter import format_document_preview
from app.services.document_checks import document_warnings
from app.services.documents import DocumentService
from app.services.reporting import ReportService
from app.services.accountant_exports import AccountantExportService


def invoice():
    return DocumentSchema(document_type='goods_invoice',external_document_number='3-0517-2',vendor='ООО Компания',date='2014-05-17',currency='RUB',total=28700,vat_total_amount=4377.97,vat_scope='document',raw_text='ТОВАРНАЯ НАКЛАДНАЯ',items=[
        dict(name='Коммутатор 8-портовый',quantity=2,price=1694.92,line_total=4000,vat_rate=18,vat_amount=610.17,line_total_without_vat=3389.83,price_includes_vat=False),
        dict(name='Коммутатор 16-портовый',quantity=1,price=4406.78,line_total=5200,vat_rate=18,vat_amount=793.22,line_total_without_vat=4406.78,price_includes_vat=False),
        dict(name='Точка доступа Wi-Fi',quantity=5,price=3305.08,line_total=19500,vat_rate=18,vat_amount=2974.58,line_total_without_vat=16525.42,price_includes_vat=False)])


class VATContractTests(unittest.TestCase):
    def test_sample_totals_are_gross_and_vat_is_not_added_twice(self):
        doc=invoice()
        self.assertEqual(document_vat_total(doc),Decimal('4377.97'))
        self.assertFalse(document_warnings(doc))
        preview=format_document_preview(doc)
        self.assertIn('В том числе НДС: 4377.97 ₽',preview)
        self.assertIn('Итого: 28700.00 ₽',preview)
        self.assertIn('НДС: 610.17 ₽',preview)
        self.assertEqual([item.quantity for item in doc.items],[2,1,5])
        self.assertEqual(DocumentItem(quantity=796).quantity,796) # Real quantities must not be blacklisted as codes.

    def test_unknown_is_not_zero_and_partial_items_do_not_invent_tax_total(self):
        doc=DocumentSchema(items=[dict(name='Item',line_total=100,vat_amount=None)])
        self.assertIsNone(document_vat_total(doc))
        self.assertIn('НДС: не определено',format_document_preview(doc))
        doc=DocumentSchema(items=[dict(name='One',vat_amount=10),dict(name='Two')])
        self.assertIsNone(document_vat_total(doc))
        doc=DocumentSchema(items=[dict(name='One',vat_amount=10),dict(name='Two',vat_amount=0)])
        self.assertEqual(document_vat_total(doc),10)
        self.assertEqual(document_vat_total(DocumentSchema(vat_scope='no_vat')),0)
        zero=DocumentSchema(vat_scope='document',vat_total_amount=0)
        self.assertIn('В том числе НДС: 0.00',format_document_preview(zero))
        self.assertNotIn('НДС: без НДС',format_document_preview(zero))

    def test_inclusive_receipt_rate_is_not_a_concatenated_number(self):
        for value,expected in (('18/118',18),('20 / 120',20),('10/110',10),('22/122',22)):
            self.assertEqual(DocumentItem(vat_rate=value).vat_rate,expected)

    def test_conflicting_header_and_lines_require_review_without_repair(self):
        doc=invoice();doc.vat_total_amount=4000
        self.assertTrue(any('НДС по позициям' in warning for warning in document_warnings(doc)))
        self.assertEqual(doc.vat_total_amount,4000)
        doc.vat_total_amount=30000
        self.assertTrue(any('превышает' in warning for warning in document_warnings(doc)))


@unittest.skipUnless(os.environ.get('RUN_DATABASE_TESTS')=='1','Disposable PostgreSQL required')
class VATPersistenceTests(DatabaseFixture):
    async def test_original_sample_vat_survives_save_and_both_exports(self):
        manager,employee=await self.finance_ids()
        doc=invoice()
        saved=await DocumentService().save_for_user(employee,self.project.id,format_document_preview(doc),doc)
        record=await self.pool.fetchrow('SELECT * FROM documents WHERE id=$1',saved)
        self.assertEqual(record['vat_total_amount'],Decimal('4377.97'))
        self.assertEqual(record['total_amount'],28700)
        items=await self.pool.fetch('SELECT * FROM document_items WHERE document_id=$1 ORDER BY line_no',saved)
        self.assertEqual([i['vat_amount'] for i in items],[Decimal('610.17'),Decimal('793.22'),Decimal('2974.58')])
        self.assertEqual([i['quantity'] for i in items],[2,1,5])
        self.assertFalse(items[0]['price_includes_vat'])
        service=ReportService();start=datetime(2000,1,1,tzinfo=timezone.utc);end=datetime(2100,1,1,tzinfo=timezone.utc)
        detail=await service.document(manager,saved)
        self.assertEqual(detail['vat_total_amount'],Decimal('4377.97'))
        self.assertEqual(detail['items'][0]['vat_rate'],18)
        report=await service.summary(manager,start,end)
        self.assertEqual(report['summary']['total_amount'],28700)
        _,data=await service.export(manager,start,end)
        book=load_workbook(BytesIO(data))
        self.assertEqual(book['Documents'].cell(2,13).value,4377.97)
        self.assertEqual(book['Items'].cell(2,8).value,610.17)
        _,data,_=await AccountantExportService().build(manager,'all_time')
        with ZipFile(BytesIO(data)) as archive:
            book=load_workbook(BytesIO(archive.read('Реестр.xlsx')))
        self.assertEqual(book['Документы'].cell(2,17).value,'4377.97')
        self.assertEqual(book['Позиции'].cell(2,8).value,'610.17')

    async def test_old_unknown_and_explicit_zero_remain_different(self):
        manager,employee=await self.finance_ids()
        first=await DocumentService().save_for_user(employee,self.project.id,'unknown',self.document)
        self.document.vat_total_amount=0;self.document.vat_scope='no_vat'
        second=await DocumentService().save_for_user(employee,self.project.id,'zero',self.document)
        self.assertIsNone(await self.pool.fetchval('SELECT vat_total_amount FROM documents WHERE id=$1',first))
        self.assertEqual(await self.pool.fetchval('SELECT vat_total_amount FROM documents WHERE id=$1',second),0)
