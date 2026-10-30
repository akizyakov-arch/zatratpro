import os
import unittest
from datetime import datetime, timezone
from io import BytesIO
from zipfile import ZipFile

from openpyxl import load_workbook
from app.schemas.document import DocumentSchema
from app.schemas.fiscal import detect_fiscalization, fiscalization_label
from app.services.json_formatter import format_document_preview
from app.services.documents import DocumentService
from app.services.reporting import ReportService
from app.services.accountant_exports import AccountantExportService
from app.services.views import ViewService
from app.services.report_formatters import format_report_document_items
from db_fixture import DatabaseFixture

FISCAL = 'Кассовый чек\nФН № 9999078900001111\nФД: 14885\nФПД = 563815926'

class FiscalEvidenceTests(unittest.TestCase):
    def test_receipt_title_tax_card_and_missing_fields_do_not_prove_fiscalization(self):
        for text in (None,'Кассовый чек ИНН 7700000001 НДС 20% ОПЛАТА КАРТОЙ', 'ФН 9999078900001111 ФД 1', 'ФД 1 ФП 123'):
            self.assertIsNone(detect_fiscalization(text))
        self.assertTrue(detect_fiscalization(FISCAL))
        self.assertTrue(detect_fiscalization('Фискальный накопитель: 9999078900001111\nФискальный документ: 14885\nФискальный признак документа: 563815926'))
        self.assertTrue(detect_fiscalization(FISCAL + '\nТовар: нефискальный регистратор'))
        self.assertTrue(detect_fiscalization('t=20261030T1200&s=100.00&fn=9999078900001111&i=14885&fp=563815926&n=1'))

    def test_negative_marker_and_incomplete_recognition(self):
        for text in ('НЕФИСКАЛЬНЫЙ ЧЕК','Документ не является фискальным','Не фискализирован',FISCAL+'\nНЕФИСКАЛЬНЫЙ ЧЕК'):
            self.assertFalse(detect_fiscalization(text))
        self.assertIsNone(detect_fiscalization('ФН ? ФД ? ФП ?'))
        self.assertIsNone(detect_fiscalization('ФН 1 ФД 2 ФП 3'))

    def test_schema_does_not_trust_unsupported_model_guess(self):
        document=DocumentSchema(raw_text='Кассовый чек',is_fiscalized=True)
        self.assertIsNone(document.is_fiscalized)
        self.assertIn('Фискализация: Не определено',format_document_preview(document))
        document=DocumentSchema(raw_text=FISCAL,is_fiscalized=False)
        self.assertTrue(document.is_fiscalized)
        self.assertIn('Фискализация: Да',format_document_preview(document))
        self.assertEqual([fiscalization_label(v) for v in (True,False,None)],['Да','Нет','Не определено'])

@unittest.skipUnless(os.environ.get('RUN_DATABASE_TESTS')=='1','Disposable PostgreSQL required')
class FiscalPersistenceTests(DatabaseFixture):
    async def test_three_states_reach_database_reports_and_both_workbooks(self):
        manager,employee=await self.finance_ids()
        records=[]
        for text in (FISCAL,'Кассовый чек. НЕФИСКАЛЬНЫЙ ЧЕК','Кассовый чек'):
            document=DocumentSchema.model_validate({**self.document.model_dump(),'raw_text':text})
            records.append(await DocumentService().save_for_user(employee,self.project.id,format_document_preview(document),document))
        rows=await self.pool.fetch('SELECT is_fiscalized FROM documents ORDER BY id')
        self.assertEqual([r['is_fiscalized'] for r in rows],[True,False,None])
        start=datetime(2000,1,1,tzinfo=timezone.utc)
        end=datetime(2100,1,1,tzinfo=timezone.utc)
        report=await ReportService().report(manager,start,end)
        self.assertEqual([d['is_fiscalized'] for d in report['documents']],[True,False,None])
        _,data=await ReportService().export(manager,start,end)
        book=load_workbook(BytesIO(data))
        self.assertEqual(book['Documents'].cell(1,12).value,'Фискализация')
        self.assertEqual([book['Documents'].cell(r,12).value for r in (2,3,4)],['Да','Нет','Не определено'])
        self.assertEqual(book['Documents'].cell(2,11).value,'RUB')
        _,data,missing=await AccountantExportService().build(manager,'all_time')
        with ZipFile(BytesIO(data)) as archive:
            book=load_workbook(BytesIO(archive.read('Реестр.xlsx')))
        self.assertEqual(book['Документы'].cell(1,16).value,'Фискализация')
        self.assertEqual([book['Документы'].cell(r,16).value for r in (2,3,4)],['Да','Нет','Не определено'])
        detail=await ViewService().get_report_document_items(manager,'all_time',records[0],project_id=self.project.id)
        self.assertTrue(detail[0].is_fiscalized)
        self.assertIn('Фискализация: Да',format_report_document_items('Документ','all_time',*detail))

    async def test_historical_documents_are_unknown_after_migration_without_backfill(self):
        _,employee=await self.finance_ids()
        await self.pool.execute('ALTER TABLE documents DROP COLUMN is_fiscalized')
        await self.pool.execute('''INSERT INTO documents(company_id,project_id,uploaded_by_user_id,document_type,source_type,raw_text,total_amount)
            VALUES($1,$2,$3,'cash_receipt','photo',$4,100)''',self.company.id,self.project.id,employee,FISCAL)
        from pathlib import Path
        await self.pool.execute((Path(__file__).resolve().parent.parent/'db/migrations/027_document_fiscalization.sql').read_text())
        row=await self.pool.fetchrow('SELECT is_fiscalized,total_amount,raw_text FROM documents')
        self.assertIsNone(row['is_fiscalized'])
        self.assertEqual(row['total_amount'],100)
        self.assertEqual(row['raw_text'],FISCAL)
