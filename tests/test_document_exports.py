import unittest
from datetime import date, datetime, timezone
from decimal import Decimal
from io import BytesIO

from openpyxl import load_workbook

from app.services.document_exports import AccountantArchiveItemRow, AccountantArchiveRow, _build_manifest


def _hyperlink_ref(cell) -> str | None:
    if cell.hyperlink is None:
        return None
    return cell.hyperlink.target or cell.hyperlink.location


class DocumentExportsManifestTests(unittest.TestCase):
    def test_manifest_contains_documents_and_positions_sheets(self) -> None:
        document = AccountantArchiveRow(
            document_id=52,
            storage_key='documents/1/52/source.jpg',
            original_filename='file_352.jpg',
            mime_type='image/jpeg',
            file_ext='.jpg',
            project_name='Кафе',
            vendor='Трансторг',
            vendor_inn='7701234567',
            document_number='N331',
            document_date=date(2026, 3, 30),
            total_amount=Decimal('550.00'),
            vat_total_amount=Decimal('91.67'),
            vat_scope='document',
            is_fiscalized=True,
            created_at=datetime(2026, 3, 30, 12, 15, tzinfo=timezone.utc),
            uploaded_by_name='Тестовый Пользователь',
            duplicate_status='none',
        )
        item = AccountantArchiveItemRow(
            document_id=52,
            line_no=1,
            name='Бизнес-ланч',
            quantity=Decimal('1.000'),
            price=Decimal('550.00'),
            line_total=Decimal('550.00'),
            vat_label='НДС 20%',
            vat_amount=Decimal('91.67'),
        )

        manifest = _build_manifest([(document, 'files/52.jpg')], [item])
        workbook = load_workbook(BytesIO(manifest))

        self.assertEqual(workbook.sheetnames, ['Документы', 'Позиции'])

        documents_sheet = workbook['Документы']
        self.assertEqual(documents_sheet['A2'].value, 52)
        self.assertEqual(documents_sheet['N2'].value, 'files/52.jpg')
        self.assertEqual(documents_sheet['O2'].value, 'Открыть файл')
        self.assertEqual(_hyperlink_ref(documents_sheet['O2']), 'files/52.jpg')

        positions_sheet = workbook['Позиции']
        self.assertEqual(positions_sheet['A2'].value, 52)
        self.assertEqual(positions_sheet['G2'].value, 1)
        self.assertEqual(positions_sheet['H2'].value, 'Бизнес-ланч')
        self.assertEqual(positions_sheet['O2'].value, 'files/52.jpg')
        self.assertEqual(positions_sheet['P2'].value, 'Открыть файл')
        self.assertEqual(_hyperlink_ref(positions_sheet['P2']), 'files/52.jpg')
        self.assertEqual(positions_sheet['Q2'].value, 'Открыть документ')
        document_link = _hyperlink_ref(positions_sheet['Q2'])
        self.assertIsNotNone(document_link)
        self.assertIn('Документы', document_link)
        self.assertTrue(document_link.endswith('A2'))

    def test_manifest_keeps_positions_sheet_when_document_has_no_items(self) -> None:
        document = AccountantArchiveRow(
            document_id=77,
            storage_key='documents/1/77/source.jpg',
            original_filename='file_377.jpg',
            mime_type='image/jpeg',
            file_ext='.jpg',
            project_name='Проект',
            vendor='Поставщик',
            vendor_inn=None,
            document_number=None,
            document_date=None,
            total_amount=Decimal('1000.00'),
            vat_total_amount=None,
            vat_scope='unknown',
            is_fiscalized=None,
            created_at=datetime(2026, 3, 30, 13, 0, tzinfo=timezone.utc),
            uploaded_by_name=None,
            duplicate_status='not_checked',
        )

        manifest = _build_manifest([(document, 'files/77.jpg')], [])
        workbook = load_workbook(BytesIO(manifest))

        self.assertEqual(workbook.sheetnames, ['Документы', 'Позиции'])
        self.assertEqual(workbook['Позиции'].max_row, 1)


if __name__ == '__main__':
    unittest.main()
