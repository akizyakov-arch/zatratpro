"""Web report regressions use isolated databases and no external APIs."""
import asyncio
import os
import unittest
from datetime import datetime, timedelta, timezone
from decimal import Decimal
from io import BytesIO
from unittest.mock import patch

from openpyxl import load_workbook

from db_fixture import DatabaseFixture
from app.services.advances import AdvanceError, AdvanceService
from app.services.business_cards import BusinessCardService
from app.services.web_reports import WebReportService


@unittest.skipUnless(os.environ.get('RUN_DATABASE_TESTS') == '1', 'PostgreSQL integration tests disabled')
class WebReportTests(DatabaseFixture):
    async def add_documents(self, count=1, *, currency='RUB', amount='10.00', project=None, user=None):
        _, employee = await self.finance_ids()
        async with self.pool.acquire() as c:
            rows = await c.fetch("""
                INSERT INTO documents(company_id,project_id,uploaded_by_user_id,document_type,
                    source_type,external_document_number,vendor,currency,total_amount)
                SELECT $1,$2,$3,'cash_receipt','photo','WEB-'||n,'Web Vendor',$4,$5
                FROM generate_series(1,$6::integer) n RETURNING id
            """, self.company.id, project or self.project.id, user or employee, currency, Decimal(amount), count)
            return [row['id'] for row in rows]

    def period(self):
        now = datetime.now(timezone.utc)
        return now - timedelta(days=1), now + timedelta(days=1)

    async def test_complete_export_exceeds_old_document_and_item_limits(self):
        manager, employee = await self.finance_ids()
        documents = await self.add_documents(63)
        async with self.pool.acquire() as c:
            await c.executemany("""INSERT INTO document_items(document_id,line_no,name,quantity,price,line_total)
                VALUES($1,$2,$3,1,1,1)""", [(doc, line, '=Untrusted item') for doc in documents for line in range(1, 7)])
        await self.add_documents(currency='USD', amount='9')
        report = await WebReportService().report(manager, *self.period())
        self.assertEqual(report['summary']['document_count'], 64)
        self.assertEqual(report['summary']['total_amount'], Decimal('630'))
        self.assertEqual(report['summary']['totals_by_currency'], {'RUB': Decimal('630'), 'USD': Decimal('9')})
        self.assertEqual(len(report['items']), 378)
        _, payload = await WebReportService().export(manager, *self.period())
        workbook = load_workbook(BytesIO(payload))
        self.assertEqual(workbook['Documents'].max_row, 65)
        self.assertEqual(workbook['Items'].max_row, 379)
        self.assertEqual(workbook['Items']['C2'].data_type, 's')
        self.assertEqual(workbook['Documents']['K65'].value, 'USD')
        self.assertEqual(set(workbook.sheetnames), {'Overview', 'Projects', 'Employees', 'Documents', 'Items', 'Duplicates', 'Money', 'BusinessCards'})
        overview = {row[0].value: row[1].value for row in workbook['Overview']}
        self.assertEqual(overview['Документов'], report['summary']['document_count'])
        self.assertEqual(overview['Сумма RUB'], report['summary']['total_amount'])

    async def test_document_filters_pagination_and_rejected_funding(self):
        manager, employee = await self.finance_ids()
        document_ids = await self.add_documents(3)
        card = await BusinessCardService().register(manager, 'Team', '1234')
        expense = await BusinessCardService().submit(employee, card, document_ids[0])
        service = WebReportService()
        rows = await service.documents(manager, payment_source='card')
        self.assertEqual(rows['total'], 1)
        self.assertEqual(rows['items'][0]['id'], document_ids[0])
        await BusinessCardService().reject(manager, expense, 'Wrong card')
        rows = await service.documents(manager, payment_source='unassigned', per_page=2, page=2, q='WEB-')
        self.assertEqual(rows['total'], 3)
        self.assertEqual(len(rows['items']), 1)
        detail = await service.document(manager, document_ids[0])
        self.assertIsNone(detail['payment_source'])
        self.assertEqual(detail['funding_history'][0]['rejection_reason'], 'Wrong card')

    async def test_document_sorting_is_applied_before_pagination_with_stable_ties(self):
        manager, _ = await self.finance_ids()
        documents = await self.add_documents(5)
        await self.pool.executemany('UPDATE documents SET total_amount=$2,vendor=$3 WHERE id=$1', [
            (documents[0], Decimal('20'), 'Beta'),
            (documents[1], Decimal('5'), 'Alpha'),
            (documents[2], Decimal('20'), 'Alpha'),
            (documents[3], Decimal('100'), 'Delta'),
            (documents[4], None, None),
        ])
        service = WebReportService()
        ascending = []
        for page in range(1, 4):
            result = await service.documents(manager, sort='amount_asc', per_page=2, page=page)
            ascending.extend(row['id'] for row in result['items'])
            self.assertEqual(result['total'], 5)
        self.assertEqual(ascending, [documents[1], documents[0], documents[2], documents[3], documents[4]])
        descending = await service.documents(manager, sort='amount_desc', per_page=2)
        self.assertEqual([row['id'] for row in descending['items']], [documents[3], documents[2]])
        vendors = await service.documents(manager, sort='name', per_page=2)
        self.assertEqual([row['id'] for row in vendors['items']], [documents[1], documents[2]])
        with self.assertRaises(AdvanceError):
            await service.documents(manager, sort='id DESC; DELETE FROM documents')

    async def test_cross_company_ids_and_employee_access_are_denied(self):
        manager, employee = await self.finance_ids()
        document = (await self.add_documents())[0]
        owner_id = await self.pool.fetchval('SELECT id FROM users WHERE telegram_id=$1', self.owner.id)
        service = WebReportService()
        for actor in (owner_id, employee):
            with self.assertRaises(AdvanceError):
                await service.documents(actor)
            with self.assertRaises(AdvanceError):
                await service.document(actor, document)
        other_company = await self.service.create_company(self.owner, 'Other')
        async with self.pool.acquire() as c:
            other_project = await c.fetchval("INSERT INTO projects(company_id,name,created_by_user_id) VALUES($1,'Other',$2) RETURNING id", other_company.id, owner_id)
            foreign_document = await c.fetchval("""INSERT INTO documents(company_id,project_id,uploaded_by_user_id,
                document_type,source_type,currency,total_amount) VALUES($1,$2,$3,'cash_receipt','photo','RUB',1) RETURNING id""", other_company.id, other_project, owner_id)
        with self.assertRaises(AdvanceError):
            await service.document(manager, foreign_document)
        with self.assertRaises(AdvanceError):
            await service.report(manager, *self.period(), project_id=other_project)

    async def test_duplicate_deletion_preserves_financial_history_even_when_rejected(self):
        manager, employee = await self.finance_ids()
        document = (await self.add_documents())[0]
        async with self.pool.acquire() as c:
            await c.execute("UPDATE documents SET duplicate_status='probable' WHERE id=$1", document)
        card = await BusinessCardService().register(manager, 'Team', '1234')
        expense = await BusinessCardService().submit(employee, card, document)
        await BusinessCardService().reject(manager, expense, 'Review')
        with self.assertRaisesRegex(AdvanceError, 'финансовой историей'):
            await WebReportService().duplicate_delete(manager, document)
        await WebReportService().duplicate_keep(manager, document)
        row = await self.pool.fetchrow('SELECT duplicate_status FROM documents WHERE id=$1', document)
        self.assertEqual(row['duplicate_status'], 'none')

    async def test_documents_and_finance_use_one_database_snapshot(self):
        manager, employee = await self.finance_ids()
        await self.add_documents()
        original = AdvanceService.report
        service = WebReportService()

        async def insert_after_document_snapshot(instance, *args, **kwargs):
            await AdvanceService().issue(manager, employee, self.project.id, '80', 'Concurrent', 'report-snapshot')
            await self.add_documents()
            return await original(instance, *args, **kwargs)

        with patch.object(AdvanceService, 'report', new=insert_after_document_snapshot):
            report = await service.report(manager, *self.period())
        self.assertEqual(report['summary']['document_count'], 1)
        self.assertEqual(report['money']['totals']['pending'], 0)
        current = await service.report(manager, *self.period())
        self.assertEqual(current['summary']['document_count'], 2)
        self.assertEqual(current['money']['totals']['pending'], Decimal('80'))

    async def test_document_and_financial_filters_apply_to_export_together(self):
        manager, employee = await self.finance_ids()
        async with self.pool.acquire() as c:
            second = await c.fetchval("INSERT INTO projects(company_id,name,created_by_user_id) VALUES($1,'Second',$2) RETURNING id", self.company.id, manager)
        await self.add_documents(2)
        await self.add_documents(3, project=second)
        await AdvanceService().issue(manager, employee, self.project.id, '10', 'One', 'one')
        await AdvanceService().issue(manager, employee, second, '30', 'Two', 'two')
        report = await WebReportService().report(manager, *self.period(), project_id=second, employee_id=employee)
        self.assertEqual(report['summary']['document_count'], 3)
        self.assertEqual(report['money']['totals']['pending'], Decimal('30'))
        _, payload = await WebReportService().export(manager, *self.period(), project_id=second, employee_id=employee)
        workbook = load_workbook(BytesIO(payload))
        self.assertEqual(workbook['Documents'].max_row, 4)
        self.assertEqual(workbook['Money']['L2'].value, 0)
        self.assertEqual(workbook['Money']['J2'].value, 30)

    async def test_disabled_manager_cannot_read_existing_document(self):
        manager, _ = await self.finance_ids()
        document = (await self.add_documents())[0]
        await self.pool.execute('UPDATE users SET is_active=FALSE WHERE id=$1', manager)
        with self.assertRaises(AdvanceError):
            await WebReportService().document(manager, document)

    async def test_overview_matches_full_report_without_loading_documents(self):
        manager, employee = await self.finance_ids()
        await self.add_documents(4)
        await self.add_documents(1, currency='USD', amount='7')
        aid = await AdvanceService().issue(manager, employee, self.project.id, '80', 'Overview', 'overview')
        await AdvanceService().receive(employee, aid)
        service = WebReportService()
        start, end = self.period()
        with patch.object(service, 'report', side_effect=AssertionError('Dashboard must not load full report')):
            result = await service.overview(manager, start, end)
        full = await service.report(manager, start, end)
        self.assertEqual(result['document_count'], full['summary']['document_count'])
        self.assertEqual(result['document_total'], full['summary']['total_amount'])
        self.assertEqual(result['totals_by_currency'], {'RUB': Decimal('40'), 'USD': Decimal('7')})
        self.assertEqual(result['received'], Decimal('80'))
        self.assertEqual(result['cash'], full['money']['cash_on_hand'])
        self.assertEqual(result['money']['totals'], full['money']['totals'])

    async def test_overview_open_projects_include_empty_and_keep_all_finance_sources_separate(self):
        manager, employee = await self.finance_ids()
        second = await self.pool.fetchval("INSERT INTO projects(company_id,name,created_by_user_id) VALUES($1,'Second',$2) RETURNING id", self.company.id, manager)
        empty = await self.pool.fetchval("INSERT INTO projects(company_id,name,created_by_user_id) VALUES($1,'Empty',$2) RETURNING id", self.company.id, manager)
        archived = await self.pool.fetchval("INSERT INTO projects(company_id,name,created_by_user_id,status) VALUES($1,'Archived',$2,'archived') RETURNING id", self.company.id, manager)
        await self.add_documents(amount='500', project=archived)
        accepted = (await self.add_documents(amount='30'))[0]
        pending = (await self.add_documents(amount='20'))[0]
        await self.add_documents(currency='USD', amount='7')
        card_pending_doc = (await self.add_documents(amount='9'))[0]
        card_accepted_doc = (await self.add_documents(amount='11', project=second))[0]
        old = (await self.add_documents(amount='99'))[0]
        start, end = self.period()
        await self.pool.execute('UPDATE documents SET created_at=$2 WHERE id=$1', old, start - timedelta(seconds=1))
        foreign_company = await self.service.create_company(self.owner, 'Foreign')
        foreign_project = await self.pool.fetchval("INSERT INTO projects(company_id,name,created_by_user_id) VALUES($1,'Foreign',$2) RETURNING id", foreign_company.id, manager)
        aid = await AdvanceService().issue(manager, employee, self.project.id, '100', 'One', 'overview-project-one')
        await AdvanceService().receive(employee, aid)
        eid = await AdvanceService().attach(employee, aid, accepted)
        await AdvanceService().approve_expense(manager, aid, eid)
        await AdvanceService().attach(employee, aid, pending)
        await AdvanceService().request_return(employee, aid, '5', 'overview-project-return')
        await AdvanceService().issue(manager, employee, second, '50', 'Two', 'overview-project-two')
        card = await BusinessCardService().register(manager, 'Projects', '4821')
        await BusinessCardService().submit(employee, card, card_pending_doc)
        ceid = await BusinessCardService().submit(employee, card, card_accepted_doc)
        await BusinessCardService().approve(manager, ceid)
        result = await WebReportService().overview(manager, start, end)
        rows = {row['project_id']: row for row in result['open_projects']}
        self.assertEqual(set(rows), {self.project.id, second, empty})
        self.assertNotIn(foreign_project, rows)
        self.assertEqual(result['project_count'], 3)
        self.assertEqual(result['document_total'], Decimal('570'))
        first = rows[self.project.id]
        self.assertEqual(first['document_count'], 4)
        self.assertEqual(first['document_total'], Decimal('59'))
        self.assertEqual(first['totals_by_currency'], {'RUB': Decimal('59'), 'USD': Decimal('7')})
        self.assertEqual((first['received'], first['spent'], first['cash'], first['debt']), (Decimal('100'), Decimal('30'), Decimal('70'), 0))
        self.assertEqual((first['pending_expenses'], first['pending_returns'], first['card_pending'], first['review_total']), (Decimal('20'), Decimal('5'), Decimal('9'), Decimal('34')))
        self.assertEqual(rows[second]['pending'], Decimal('50'))
        self.assertEqual(rows[second]['card_spent'], Decimal('11'))
        self.assertEqual(rows[second]['cash'], 0)
        for key in ('document_count', 'document_total', 'received', 'cash', 'debt', 'review_total', 'card_spent', 'pending'):
            self.assertEqual(rows[empty][key], 0, key)

    async def test_project_cash_and_debt_do_not_offset_people_or_other_projects(self):
        manager, employee = await self.finance_ids()
        colleague = self.user(1004)
        code = await self.service.create_invite(self.manager, 'employee')
        await self.service.join_company(colleague, code)
        colleague_id = await self.pool.fetchval('SELECT id FROM users WHERE telegram_id=$1', colleague.id)
        second = await self.pool.fetchval("INSERT INTO projects(company_id,name,created_by_user_id) VALUES($1,'Second',$2) RETURNING id", self.company.id, manager)
        a = await AdvanceService().issue(manager, employee, self.project.id, '100', 'One', 'project-debt')
        await AdvanceService().receive(employee, a)
        doc = (await self.add_documents(amount='150'))[0]
        e = await AdvanceService().attach(employee, a, doc)
        await AdvanceService().approve_expense(manager, a, e)
        b = await AdvanceService().issue(manager, employee, second, '50', 'Two', 'project-cash')
        await AdvanceService().receive(employee, b)
        c = await AdvanceService().issue(manager, colleague_id, self.project.id, '30', 'Colleague', 'project-colleague')
        await AdvanceService().receive(colleague_id, c)
        result = await WebReportService().overview(manager, *self.period())
        rows = {row['project_id']: row for row in result['open_projects']}
        self.assertEqual(rows[self.project.id]['cash'], Decimal('30'))
        self.assertEqual(rows[self.project.id]['debt'], Decimal('50'))
        self.assertEqual(rows[second]['cash'], Decimal('50'))
        self.assertEqual(rows[second]['debt'], 0)
        self.assertEqual(result['cash'], Decimal('30'))
        self.assertEqual(result['debt'], 0)

    async def test_duplicate_delete_waits_on_company_before_locking_document(self):
        manager, employee = await self.finance_ids()
        document = (await self.add_documents())[0]
        await self.pool.execute("UPDATE documents SET duplicate_status='probable' WHERE id=$1", document)
        card = await BusinessCardService().register(manager, 'Lock order', '9234')
        task = None
        submit = None
        try:
            async with self.pool.acquire() as blocker:
                async with blocker.transaction():
                    await blocker.execute('SELECT id FROM companies WHERE id=$1 FOR UPDATE', self.company.id)
                    task = asyncio.create_task(WebReportService().duplicate_delete(manager, document))
                    async with self.pool.acquire() as observer:
                        for _ in range(100):
                            waiting = await observer.fetchval("""SELECT 1 FROM pg_stat_activity
                                WHERE datname=current_database() AND wait_event_type='Lock'
                                AND query LIKE 'SELECT id FROM companies WHERE id=$1 FOR UPDATE%'
                                LIMIT 1""")
                            if waiting:
                                break
                            await asyncio.sleep(.01)
                        else:
                            self.fail('Duplicate deletion did not wait on the company lock')
                        # Before the common company lock is acquired, deletion must
                        # not hold the document needed by a financing operation.
                        async with observer.transaction():
                            self.assertEqual(await observer.fetchval(
                                'SELECT id FROM documents WHERE id=$1 FOR UPDATE NOWAIT', document), document)
                    submit = asyncio.create_task(BusinessCardService().submit(employee, card, document))
            outcomes = await asyncio.wait_for(asyncio.gather(task, submit, return_exceptions=True), 5)
            self.assertEqual(sum(not isinstance(value, Exception) for value in outcomes), 1)
            self.assertEqual(sum(isinstance(value, AdvanceError) for value in outcomes), 1)
        finally:
            pending = [task for task in (task, submit) if task is not None and not task.done()]
            for task in pending:
                task.cancel()
            await asyncio.gather(*pending, return_exceptions=True)
