"""Cross-channel HTTP workflows use disposable databases and no external APIs."""
import os
import unittest
from datetime import datetime, time, timedelta, timezone
from html.parser import HTMLParser
from io import BytesIO
from unittest.mock import patch

import httpx
from openpyxl import load_workbook

from db_fixture import DatabaseFixture
from app.services.advances import AdvanceService
from app.services.business_cards import BusinessCardService
from app.adapters.telegram.documents import DocumentService
from app.adapters.telegram.projects import ProjectService
from app.services.web_auth import WebAuthService
from app.web.app import create_app


class AssetParser(HTMLParser):
    def __init__(self):
        super().__init__()
        self.paths = []

    def handle_starttag(self, tag, attrs):
        attrs = dict(attrs)
        if tag == 'script' and attrs.get('src'):
            self.paths.append(attrs['src'])
        elif tag == 'link' and attrs.get('rel') == 'stylesheet':
            self.paths.append(attrs['href'])


@unittest.skipUnless(os.environ.get('RUN_DATABASE_TESTS') == '1', 'Requires disposable PostgreSQL databases')
class WebAPITests(DatabaseFixture):
    async def asyncSetUp(self):
        await super().asyncSetUp()
        env = patch.dict(os.environ, {'WEB_ORIGIN': 'https://testserver', 'WEB_COOKIE_SECURE': 'true'})
        env.start()
        self.addCleanup(env.stop)
        self.client = httpx.AsyncClient(transport=httpx.ASGITransport(app=create_app()), base_url='https://testserver')
        self.addAsyncCleanup(self.client.aclose)
        self.csrf = None

    async def login(self, user=None):
        self.client.cookies.clear()
        response = await self.client.post('/api/auth/start', headers={'origin': 'https://testserver'})
        self.assertEqual(response.status_code, 200, response.text)
        request = response.json()
        code = request['bot_url'].split('web_', 1)[1]
        service = WebAuthService()
        telegram_user = user or self.manager
        await service.present(telegram_user.id, code)
        await service.decide(telegram_user.id, request['id'], True)
        response = await self.client.get('/api/auth/status')
        self.assertEqual(response.status_code, 200, response.text)
        self.assertEqual(response.json()['status'], 'authenticated')
        self.csrf = response.json()['csrf_token']
        self.assertNotIn('session_token', response.json())
        self.assertIn('Secure', response.headers['set-cookie'])
        return response.json()['user']

    @property
    def headers(self):
        return {'origin': 'https://testserver', 'x-csrf-token': self.csrf}

    async def issue(self, *, amount='200.10', key='http-cross-channel'):
        _, employee = await self.finance_ids()
        response = await self.client.post('/api/advances', headers=self.headers, json={
            'employee_id': employee, 'project_id': self.project.id, 'amount': amount,
            'purpose': 'Материалы на объект', 'request_key': key,
        })
        self.assertEqual(response.status_code, 201, response.text)
        return response.json()['id']

    async def save_document(self, *, user=None, project=None):
        return await DocumentService().save_document(
            user or self.employee, project or self.project, 'Receipt preview', self.document,
        )

    async def test_web_issue_bot_receipt_automatic_cost_report_and_excel_agree(self):
        user = await self.login()
        manager, employee = await self.finance_ids()
        self.assertEqual(user['id'], manager)
        aid = await self.issue()
        self.assertEqual(await self.issue(), aid, 'Retry must retain the same operation')
        pending = (await self.client.get('/api/overview?period=all_time')).json()
        self.assertEqual(pending['pending'], '200.10')
        self.assertEqual(pending['cash'], '0')
        await AdvanceService().receive(employee, aid)
        document = await self.save_document()
        eid = await AdvanceService().attach(employee, aid, document)
        review = await self.client.get('/api/reviews')
        self.assertEqual(review.status_code, 200, review.text)
        self.assertEqual(review.json()['total'], 0)
        self.assertEqual(str((await AdvanceService().card(employee, aid))['balance']), '100.10')
        # Compatibility route for old buttons is an idempotent no-op.
        response = await self.client.post(f'/api/advances/{aid}/expenses/{eid}/approve', headers=self.headers)
        self.assertEqual(response.status_code, 200, response.text)
        self.assertEqual((await self.client.get('/api/reviews')).json()['total'], 0)
        bot_card = await AdvanceService().card(employee, aid)
        self.assertEqual(str(bot_card['balance']), '100.10')
        card = await self.client.get(f'/api/advances/{aid}')
        self.assertEqual(card.status_code, 200, card.text)
        self.assertEqual(card.json()['balance'], '100.10')
        report = await self.client.get('/api/reports?period=all_time')
        self.assertEqual(report.status_code, 200, report.text)
        data = report.json()
        self.assertEqual(data['summary']['total_amount'], '100.00')
        self.assertEqual(data['money']['totals']['received'], '200.10')
        self.assertEqual(data['money']['totals']['spent'], '100.00')
        self.assertEqual(data['money']['totals']['closing'], '100.10')
        exported = await self.client.get('/api/reports/export?period=all_time')
        self.assertEqual(exported.status_code, 200, exported.text[:100] if exported.status_code != 200 else '')
        workbook = load_workbook(BytesIO(exported.content))
        self.assertAlmostEqual(workbook['Money']['I2'].value, 100.10)
        self.assertEqual(workbook['Documents'].max_row, 2)
        self.assertIn('attachment;', exported.headers['content-disposition'])
        self.assertEqual(exported.headers['cache-control'], 'no-store')

    async def test_review_queue_only_requires_cash_return_confirmation(self):
        await self.login()
        manager, employee = await self.finance_ids()
        aid = await self.issue(amount='300')
        await AdvanceService().receive(employee, aid)
        rid = await AdvanceService().request_return(employee, aid, '20', 'queue-return')
        card_response = await self.client.post('/api/cards', headers=self.headers, json={'name': 'Закупки', 'last4': '0123'})
        self.assertEqual(card_response.status_code, 201, card_response.text)
        card_id = card_response.json()['id']
        ceid = await BusinessCardService().submit(employee, card_id, await self.save_document())
        page = await self.client.get('/api/reviews?per_page=1')
        self.assertEqual(page.json()['total'], 1)
        self.assertEqual(len(page.json()['items']), 1)
        second = await self.client.get('/api/reviews?per_page=1&page=2')
        self.assertEqual(page.json()['items'][0]['type'], 'return')
        self.assertEqual(second.json()['items'], [])
        rejected = await self.client.post(f'/api/card-expenses/{ceid}/reject', headers=self.headers, json={'reason': 'Оплата другим способом'})
        self.assertEqual(rejected.status_code, 400, rejected.text)
        approved = await self.client.post(f'/api/advances/{aid}/returns/{rid}/approve', headers=self.headers)
        self.assertEqual(approved.status_code, 200, approved.text)
        self.assertEqual((await self.client.get('/api/reviews')).json()['total'], 0)
        rows = (await self.client.get('/api/card-expenses?status=approved')).json()
        self.assertEqual(rows['total'], 1)
        self.assertIsNotNone(rows['items'][0]['approved_at'])
        self.assertEqual(str((await AdvanceService().card(employee, aid))['balance']), '280.00')

    async def test_duplicate_warning_and_list_include_archives_outside_selected_period(self):
        await self.login()
        original = await self.save_document()
        copy = await self.save_document()
        await self.pool.execute("UPDATE projects SET status='archived',archived_at=clock_timestamp() WHERE id=$1", self.project.id)
        await self.pool.execute("UPDATE documents SET created_at='2001-01-01' WHERE id=ANY($1::bigint[])", [original, copy])
        result = await self.client.get('/api/overview?period=month')
        self.assertEqual(result.status_code, 200, result.text)
        self.assertEqual(result.json()['document_count'], 0)
        self.assertEqual(result.json()['duplicate_review'], {
            'total': 1, 'exact': 1, 'probable': 0, 'involving_archived_projects': 1,
        })
        rows = await self.client.get('/api/documents?period=all_time&duplicate_status=flagged')
        self.assertEqual(rows.status_code, 200, rows.text)
        self.assertEqual(rows.json()['total'], 1)
        self.assertEqual(rows.json()['items'][0]['id'], copy)
        self.assertEqual(rows.json()['items'][0]['project_status'], 'archived')

    async def test_advance_purpose_can_be_omitted_empty_or_null(self):
        await self.login()
        _, employee = await self.finance_ids()
        for index, extra in enumerate(({}, {'purpose': ''}, {'purpose': None}, {'purpose': '   '})):
            response = await self.client.post('/api/advances', headers=self.headers, json={
                'employee_id': employee, 'project_id': self.project.id, 'amount': '50',
                'request_key': f'optional-purpose-{index}', **extra,
            })
            self.assertEqual(response.status_code, 201, response.text)
            card = await self.client.get(f"/api/advances/{response.json()['id']}")
            self.assertEqual(card.json()['purpose'], '')

    async def test_api_requires_session_same_origin_csrf_and_current_role(self):
        for endpoint in ('/api/projects', '/api/reviews', '/api/reports/export', '/api/owner/stats'):
            self.assertEqual((await self.client.get(endpoint)).status_code, 401)
        manager = await self.login()
        self.assertEqual((await self.client.post('/api/projects', json={'name': 'No CSRF'})).status_code, 403)
        foreign_origin = {'origin': 'https://evil.test', 'x-csrf-token': self.csrf}
        self.assertEqual((await self.client.post('/api/projects', headers=foreign_origin, json={'name': 'Bad origin'})).status_code, 403)
        self.assertEqual((await self.client.get('/api/owner/stats')).status_code, 403)
        self.assertEqual((await self.client.post('/api/owner/companies', headers=self.headers, json={'name': 'No owner rights'})).status_code, 403)
        await self.pool.execute('UPDATE users SET is_active=false WHERE id=$1', manager['id'])
        self.assertEqual((await self.client.get('/api/projects')).status_code, 401)

    async def test_owner_can_manage_companies_but_cannot_perform_finance(self):
        await self.login(self.owner)
        self.assertEqual((await self.client.get('/api/owner/stats')).status_code, 200)
        created = await self.client.post('/api/owner/companies', headers=self.headers, json={'name': 'Новая компания'})
        self.assertEqual(created.status_code, 201, created.text)
        company_id = created.json()['id']
        invite = await self.client.post(f'/api/owner/companies/{company_id}/invite', headers=self.headers)
        self.assertEqual(invite.status_code, 200, invite.text)
        self.assertTrue(invite.json()['code'])
        self.assertEqual((await self.client.get('/api/advances')).status_code, 403)
        _, employee = await self.finance_ids()
        response = await self.client.post('/api/advances', headers=self.headers, json={
            'employee_id': employee, 'project_id': self.project.id, 'amount': '1',
            'purpose': 'No owner financial privileges', 'request_key': 'owner-forbidden',
        })
        self.assertEqual(response.status_code, 403)
        self.assertEqual(await self.pool.fetchval('SELECT count(*) FROM advances'), 0)

    async def test_foreign_ids_and_invalid_input_do_not_read_or_mutate(self):
        await self.login()
        other = await self.service.create_company(self.owner, 'Foreign company')
        other_manager = self.user(1004)
        code = await self.service.create_initial_manager_invite(self.owner, other.id)
        await self.service.join_company(other_manager, code)
        foreign_project = await ProjectService().create_project(other_manager.id, 'Confidential project')
        foreign_document = await self.save_document(user=other_manager, project=foreign_project)
        for endpoint in (f'/api/documents/{foreign_document}', f'/api/projects/{foreign_project.id}', f'/api/reports?project_id={foreign_project.id}'):
            response = await self.client.get(endpoint)
            self.assertEqual(response.status_code, 400, response.text)
            self.assertNotIn('Confidential', response.text)
        response = await self.client.patch(f'/api/projects/{foreign_project.id}', headers=self.headers, json={'name': 'Changed'})
        self.assertEqual(response.status_code, 400)
        self.assertEqual(await self.pool.fetchval('SELECT name FROM projects WHERE id=$1', foreign_project.id), 'Confidential project')
        self.assertEqual((await self.client.get('/api/advances?page=0')).status_code, 422)
        self.assertEqual((await self.client.get('/api/documents?per_page=101')).status_code, 422)
        self.assertEqual((await self.client.get('/api/reports?start=2026-10-25&end=2026-10-24')).status_code, 422)
        response = await self.client.post('/api/projects', headers=self.headers, json={'name': '  ', 'company_id': other.id})
        self.assertEqual(response.status_code, 422)
        _, employee = await self.finance_ids()
        response = await self.client.post('/api/advances', headers=self.headers, json={
            'employee_id': employee, 'project_id': self.project.id, 'amount': 'NaN',
            'purpose': 'Invalid money', 'request_key': 'invalid-money',
        })
        self.assertEqual(response.status_code, 400)
        self.assertEqual(await self.pool.fetchval('SELECT count(*) FROM advances'), 0)

    async def test_static_shell_is_public_but_contains_no_company_data(self):
        response = await self.client.get('/')
        self.assertEqual(response.status_code, 200, response.text)
        self.assertIn('FIINOPS', response.text)
        self.assertNotIn(self.company.name, response.text)
        self.assertNotIn('TEST_ONLY', response.text)
        self.assertEqual(response.headers['x-frame-options'], 'DENY')
        self.assertEqual(response.headers['x-content-type-options'], 'nosniff')
        self.assertIn("script-src 'self'", response.headers['content-security-policy'])
        self.assertEqual((await self.client.get('/api/company')).status_code, 401)
        assets = AssetParser()
        assets.feed(response.text)
        self.assertGreaterEqual(len(assets.paths), 2)
        for path in assets.paths:
            self.assertTrue(path.startswith('/assets/'), path)
            self.assertEqual((await self.client.get(path)).status_code, 200, path)

    async def test_date_edges_unknown_routes_and_untrusted_host_fail_cleanly(self):
        await self.login()
        for path in (
            '/api/advances?end=9999-12-31',
            '/api/card-expenses?end=9999-12-31',
            '/api/reports?end=9999-12-31',
            '/api/advances?start=2026-10-25&end=2026-10-24',
            '/api/card-expenses?start=2026-10-25&end=2026-10-24',
        ):
            response = await self.client.get(path)
            self.assertEqual(response.status_code, 422, path + ': ' + response.text)
            self.assertNotIn('Traceback', response.text)
        response = await self.client.get('/api/unknown-resource')
        self.assertEqual(response.status_code, 404)
        self.assertNotIn('Traceback', response.text)
        response = await self.client.get('/api/company', headers={'host': 'untrusted.example'})
        self.assertEqual(response.status_code, 400)

    async def test_ledger_periods_filter_week_and_include_custom_end_day(self):
        await self.login()
        manager, employee = await self.finance_ids()
        old_advance = await self.issue(amount='10', key='period-old')
        current_advance = await self.issue(amount='20', key='period-current')
        card = await BusinessCardService().register(manager, 'Period card', '4567')
        old_card_expense = await BusinessCardService().submit(employee, card, await self.save_document())
        current_card_expense = await BusinessCardService().submit(employee, card, await self.save_document())
        old_day = (datetime.now(timezone.utc) - timedelta(days=14)).date()
        old_timestamp = datetime.combine(old_day, time(23, 59, 59), timezone.utc)
        await self.pool.execute('UPDATE advances SET created_at=$2 WHERE id=$1', old_advance, old_timestamp)
        await self.pool.execute('UPDATE card_expenses SET created_at=$2 WHERE id=$1', old_card_expense, old_timestamp)
        for path, old_id, current_id in (
            ('/api/advances', old_advance, current_advance),
            ('/api/card-expenses', old_card_expense, current_card_expense),
        ):
            entire_history = await self.client.get(path)
            self.assertEqual(entire_history.status_code, 200, entire_history.text)
            self.assertEqual(entire_history.json()['total'], 2)
            week = await self.client.get(path + '?period=week')
            self.assertEqual(week.status_code, 200, week.text)
            self.assertEqual([row['id'] for row in week.json()['items']], [current_id])
            custom = await self.client.get(path, params={'period': 'custom', 'start': str(old_day), 'end': str(old_day)})
            self.assertEqual(custom.status_code, 200, custom.text)
            self.assertEqual([row['id'] for row in custom.json()['items']], [old_id])
            invalid = await self.client.get(path + '?period=unknown-period')
            self.assertEqual(invalid.status_code, 422, invalid.text)

    async def test_max_browser_login_uses_secure_cookie_and_same_permissions(self):
        manager, _ = await self.finance_ids()
        await self.pool.execute("INSERT INTO external_accounts(user_id,platform,external_id) VALUES($1,'max','9001')", manager)
        with patch('app.web.max_login.login_link', return_value='https://max.ru/test_bot?start=web_test'):
            response = await self.client.post('/api/auth/start?platform=max', headers={'origin': 'https://testserver'})
        self.assertEqual(response.status_code, 200, response.text)
        data = response.json()
        self.assertEqual(data['bot_url'], 'https://max.ru/test_bot?start=web_test')
        self.assertIn('Secure', response.headers['set-cookie'])
        self.assertIn('HttpOnly', response.headers['set-cookie'])
        await WebAuthService().present(9001, data['command'].split()[1], platform='max')
        await WebAuthService().decide(9001, data['id'], True, platform='max')
        response = await self.client.get('/api/auth/status')
        self.assertEqual(response.json()['status'], 'authenticated')
        self.assertEqual(response.json()['user']['id'], manager)
        self.assertNotIn('session_token', response.json())
        self.assertEqual((await self.client.get('/api/auth/me')).status_code, 200)
        bad = await self.client.post('/api/auth/start?platform=other', headers={'origin': 'https://testserver'})
        self.assertEqual(bad.status_code, 422)
